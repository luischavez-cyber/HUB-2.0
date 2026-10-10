"""Utilidades de formato Word para el informe completo (python-docx)."""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

RE_PENDIENTE = re.compile(r"\[(?:COMPLETAR|DATO PENDIENTE)[^\]]*\]")
GRIS_ENCABEZADO = "D9D9D9"


# ------------------------------------------------------------------ documento base
def documento(plantilla: Path | bytes | None = None) -> Any:
    """Documento nuevo o basado en la plantilla de CAEPSA (se conservan estilos, encabezados y pies)."""
    if plantilla is None:
        doc = Document()
        for sec in doc.sections:
            sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # carta
            sec.left_margin = sec.right_margin = Cm(2.5)
            sec.top_margin = sec.bottom_margin = Cm(2.5)
    else:
        doc = Document(io.BytesIO(plantilla) if isinstance(plantilla, bytes) else str(plantilla))
        _limpiar_cuerpo(doc)
    asegurar_estilos(doc)
    return doc


def _limpiar_cuerpo(doc: Any) -> None:
    """Quita el contenido de la plantilla conservando la última definición de sección."""
    cuerpo = doc.element.body
    for el in list(cuerpo):
        if el.tag != qn("w:sectPr"):
            cuerpo.remove(el)


def _tiene_estilo(doc: Any, nombre: str) -> bool:
    try:
        doc.styles[nombre]
        return True
    except KeyError:
        return False


def asegurar_estilos(doc: Any) -> None:
    """Crea los estilos que use el informe y que la plantilla no defina (Word solo guarda los estilos usados)."""
    for nivel, tam in ((1, 14), (2, 12), (3, 11)):
        nombre = f"Heading {nivel}"
        if not _tiene_estilo(doc, nombre):
            st = doc.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = doc.styles["Normal"]
            st.font.bold = True
            st.font.size = Pt(tam)
            st.paragraph_format.space_before = Pt(12)
            st.paragraph_format.keep_with_next = True
            ppr = st.element.get_or_add_pPr()
            lvl = OxmlElement("w:outlineLvl")  # necesario para que el índice lo reconozca
            lvl.set(qn("w:val"), str(nivel - 1))
            ppr.append(lvl)
    if not _tiene_estilo(doc, "Table Grid"):
        st = doc.styles.add_style("Table Grid", WD_STYLE_TYPE.TABLE)
        tblpr = st.element.find(qn("w:tblPr"))
        if tblpr is None:
            tblpr = OxmlElement("w:tblPr")
            st.element.append(tblpr)
        bordes = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{lado}")
            for k, v in (("val", "single"), ("sz", "4"), ("space", "0"), ("color", "000000")):
                b.set(qn(f"w:{k}"), v)
            bordes.append(b)
        tblpr.append(bordes)


def actualizar_campos_al_abrir(doc: Any) -> None:
    """Word pedirá actualizar el índice y la paginación al abrir el documento."""
    ajustes = doc.settings.element
    if ajustes.find(qn("w:updateFields")) is None:
        u = OxmlElement("w:updateFields")
        u.set(qn("w:val"), "true")
        ajustes.append(u)


# ------------------------------------------------------------------ campos
def _campo(parrafo: Any, instruccion: str, marcador: str = "") -> None:
    def fld(tipo: str) -> Any:
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), tipo)
        return e

    r1 = parrafo.add_run()
    r1._r.append(fld("begin"))
    r2 = parrafo.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruccion} "
    r2._r.append(instr)
    r3 = parrafo.add_run()
    r3._r.append(fld("separate"))
    parrafo.add_run(marcador)
    r4 = parrafo.add_run()
    r4._r.append(fld("end"))


MARCADOR_INDICE = "Actualice el índice: clic derecho sobre este texto → «Actualizar campo»."


def indice(doc: Any) -> None:
    # Título como párrafo en negritas (no «Heading») para que el índice no se incluya a sí mismo.
    par_t = doc.add_paragraph()
    r = par_t.add_run("ÍNDICE")
    r.bold = True
    r.font.size = Pt(14)
    par = doc.add_paragraph()
    _campo(par, 'TOC \\o "1-2" \\h \\z \\u', MARCADOR_INDICE)


def precargar_indice(doc: Any) -> None:
    """Rellena el resultado del campo con los títulos (sin páginas); Word lo actualiza al abrir."""
    from docx.enum.text import WD_BREAK

    titulos = [(p.style.name, p.text) for p in doc.paragraphs
               if p.style is not None and p.style.name in ("Heading 1", "Heading 2") and p.text.strip()]
    for par in doc.paragraphs:
        for run in par.runs:
            if run.text == MARCADOR_INDICE:
                run.text = ""
                for i, (estilo, texto) in enumerate(titulos):
                    if i:
                        run.add_break(WD_BREAK.LINE)
                    run.add_text(("      " if estilo == "Heading 2" else "") + texto)
                return


def pie_paginado(doc: Any, texto: str) -> None:
    """Pie «texto · Página X de Y» si la plantilla no trae paginación propia."""
    # Solo la primera sección: las demás heredan su pie (vinculadas a la anterior).
    for sec in doc.sections[:1]:
        pie = sec.footer
        if any("PAGE" in (i.text or "") for i in pie._element.iter(qn("w:instrText"))):
            continue
        par = pie.paragraphs[0] if pie.paragraphs else pie.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(f"{texto} · Página ")
        r.font.size = Pt(8)
        _campo(par, "PAGE", "1")
        par.add_run(" de ").font.size = Pt(8)
        _campo(par, "NUMPAGES", "1")


def encabezado(doc: Any, texto: str) -> None:
    """Encabezado con el nombre del proyecto si la plantilla no trae uno."""
    for sec in doc.sections[:1]:
        enc = sec.header
        if "".join(enc._element.itertext()).strip() or enc._element.findall(".//" + qn("w:drawing")):
            continue
        par = enc.paragraphs[0] if enc.paragraphs else enc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = par.add_run(texto)
        r.font.size = Pt(8)
        r.italic = True


# ------------------------------------------------------------------ secciones de página
def salto_pagina(doc: Any) -> None:
    from docx.enum.text import WD_BREAK

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def seccion_horizontal(doc: Any) -> None:
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    if sec.orientation != WD_ORIENT.LANDSCAPE:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width


def seccion_vertical(doc: Any) -> None:
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    if sec.orientation != WD_ORIENT.PORTRAIT:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = sec.page_height, sec.page_width


# ------------------------------------------------------------------ texto y tablas
def parrafo(doc_o_celda: Any, texto: str, negrita: bool = False, tam: float | None = None,
            alinear: Any = WD_ALIGN_PARAGRAPH.JUSTIFY) -> Any:
    """Párrafo con los marcadores [COMPLETAR…]/[DATO PENDIENTE…] resaltados en amarillo."""
    par = doc_o_celda.add_paragraph()
    par.alignment = alinear
    pos = 0
    for m in RE_PENDIENTE.finditer(texto):
        if m.start() > pos:
            _run(par, texto[pos:m.start()], negrita, tam)
        _run(par, m.group(), negrita, tam).font.highlight_color = WD_COLOR_INDEX.YELLOW
        pos = m.end()
    if pos < len(texto):
        _run(par, texto[pos:], negrita, tam)
    return par


def _run(par: Any, texto: str, negrita: bool, tam: float | None) -> Any:
    r = par.add_run(texto)
    r.bold = negrita
    if tam:
        r.font.size = Pt(tam)
    return r


def titulo_tabla(doc: Any, texto: str) -> None:
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    r = par.add_run(texto)
    r.bold = True
    r.font.size = Pt(10)


def tabla(doc: Any, encabezados: list[str], filas: list[list[Any]], tam: float = 8,
          anchos_cm: list[float] | None = None) -> Any:
    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, encabezados):
        c.text = ""
        _run(c.paragraphs[0], h, True, tam)
        tc = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), GRIS_ENCABEZADO)
        tc.append(shd)
    for fila in filas:
        celdas = t.add_row().cells
        for c, v in zip(celdas, fila):
            c.text = ""
            texto = "—" if v is None or v == "" else str(v)
            par = c.paragraphs[0]
            pos = 0
            for m in RE_PENDIENTE.finditer(texto):
                if m.start() > pos:
                    _run(par, texto[pos:m.start()], False, tam)
                _run(par, m.group(), False, tam).font.highlight_color = WD_COLOR_INDEX.YELLOW
                pos = m.end()
            if pos < len(texto):
                _run(par, texto[pos:], False, tam)
    if anchos_cm:
        for fila in t.rows:
            for c, w in zip(fila.cells, anchos_cm):
                c.width = Cm(w)
    return t


def contar_pendientes(doc: Any) -> list[str]:
    """Marcadores [COMPLETAR…]/[DATO PENDIENTE…] presentes en el cuerpo del documento."""
    texto = "\n".join(p.text for p in doc.paragraphs)
    texto += "\n" + "\n".join(c.text for t in doc.tables for c in t._cells)
    return RE_PENDIENTE.findall(texto)
