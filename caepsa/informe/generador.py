"""Módulo 7 — Generación del Informe de Caracterización completo (Word).

Todo el contenido proviene de los módulos anteriores. Lo que la plataforma no puede sustentar con datos se
inserta como «[COMPLETAR: …]» resaltado, para que el responsable técnico lo redacte; nunca se inventa.
"""
from __future__ import annotations

import io
from datetime import date
from pathlib import Path
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..dispersion import capitulo as cap5
from ..dispersion import medio_fisico as mf
from ..exportacion.tablas import escribir_tabla_nom138, formatear_valor
from ..fotografia.registro import diagnostico as diag_fotos
from ..fotografia.seccion import escribir_seccion
from ..normativa.evaluacion import CUMPLE, EXCEDE, INDETERMINADO, evaluar, resumen
from ..normativa.nom138 import (NORMA, PARAMETROS, PUBLICACION_DOF, REFERENCIAS, TITULO, UNIDAD, lmp_aplicable,
                                parametros_requeridos)
from ..proyectos import CAMPOS_FICHA, Proyecto, campos_faltantes
from ..revision.registro import hallazgos_de, listar_revisiones
from . import formato as fmt

CONSULTORA = "Consultoría Ambiental Estudios y Proyectos S.A. de C.V. (CAEPSA)"
MESES = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre",
         "noviembre", "diciembre")
CAPITULOS = (
    "1. ANTECEDENTES", "2. MARCO LEGAL", "3. METODOLOGÍA DE MUESTREO", "4. RESULTADOS Y EVALUACIÓN NORMATIVA",
    "5. ANÁLISIS DE DISPERSIÓN DE HIDROCARBUROS", "6. IDENTIFICACIÓN DE LA PLUMA CONTAMINANTE",
    "7. EVALUACIÓN DE DAÑOS", "8. CONCLUSIONES", "9. PLAN DE SANEAMIENTO", "10. MEMORIA FOTOGRÁFICA", "ANEXOS",
)


def completar(texto: str) -> str:
    return f"[COMPLETAR: {texto}]"


def _lista(xs: list[str]) -> str:
    xs = [x for x in xs if x]
    return xs[0] if len(xs) == 1 else ", ".join(xs[:-1]) + " y " + xs[-1] if xs else ""


class _Contexto:
    """Datos del proyecto que usan todos los capítulos (se calculan una sola vez)."""

    def __init__(self, p: Proyecto):
        self.p = p
        self.ficha = p.ficha()
        self.usos = p.usos_suelo()
        self.producto = self.ficha.get("producto")
        self.filas = p.tabla_resultados()
        hay = bool(self.usos) and any(f["parametro"] for f in self.filas)
        orden = {z: i for i, z in enumerate(p.zonas())}

        def clave_zona(zona: str | None, *resto: Any) -> tuple[Any, ...]:
            return (orden.get(zona or "", len(orden)), zona or "", *resto)

        self.evaluadas = sorted(evaluar(self.filas, self.usos, self.producto) if hay else [],
                                key=lambda m: clave_zona(m.zona, m.punto or "", m.muestra))
        self.resumen = resumen(self.evaluadas) if hay else None
        self.puntos = sorted([dict(r) for r in p.con.execute("SELECT * FROM puntos")],
                             key=lambda x: clave_zona(x["zona"], x["clave"]))
        self.medio = mf.leer(p)
        versiones = cap5.listar(p)
        self.cap5 = versiones[0] if versiones else cap5.registro_al_vuelo(p)
        self.cap5_guardado = bool(versiones)
        self.a = self.cap5["metricas"]

    def f(self, campo: str, etiqueta: str | None = None) -> str:
        v = self.ficha.get(campo)
        return str(v) if v else completar(etiqueta or campo)

    def usos_txt(self) -> str:
        return " / ".join(u.etiqueta.lower() for u in self.usos) if self.usos else completar("uso de suelo")

    def mf(self, clave: str) -> Any:
        return mf.valor(self.medio, clave)


# ------------------------------------------------------------------ capítulos
def _portada(doc: Any, c: _Contexto, hoy: date) -> None:
    for _ in range(5):
        doc.add_paragraph()
    for texto, tam in (("INFORME DE CARACTERIZACIÓN", 22),
                       ("Suelo afectado por derrame de hidrocarburos", 14)):
        fmt.parrafo(doc, texto, negrita=True, tam=tam, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    fmt.parrafo(doc, c.f("nombre", "nombre del proyecto"), negrita=True, tam=14, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    lugar = ", ".join(x for x in (c.ficha.get("ubicacion"), c.ficha.get("municipio"), c.ficha.get("estado")) if x)
    fmt.parrafo(doc, lugar or completar("ubicación del sitio"), tam=12, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    for _ in range(4):
        doc.add_paragraph()
    for etq, val in (("Cliente", c.f("cliente", "cliente")), ("Elaboró", CONSULTORA),
                     ("Responsable técnico", c.f("responsable_tecnico", "responsable técnico")),
                     ("Número de expediente", c.f("expediente", "número de expediente"))):
        fmt.parrafo(doc, f"{etq}: {val}", tam=11, alinear=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    fmt.parrafo(doc, f"{MESES[hoy.month - 1].capitalize()} de {hoy.year}", tam=11, alinear=WD_ALIGN_PARAGRAPH.CENTER)


def _capitulo(doc: Any, titulo: str) -> None:
    fmt.salto_pagina(doc)
    doc.add_heading(titulo, level=1)


def _antecedentes(doc: Any, c: _Contexto) -> None:
    _capitulo(doc, CAPITULOS[0])
    doc.add_heading("1.1 Datos generales del proyecto", level=2)
    filas = [[campo.etiqueta, c.f(campo.clave, campo.etiqueta.lower())]
             for campo in CAMPOS_FICHA if campo.clave not in ("zonas", "puntos_esperados", "muestras_esperadas",
                                                              "usos_suelo", "utm_zona")]
    filas.insert(7, ["Uso de suelo", c.usos_txt()])
    filas.append(["Elaboró", CONSULTORA])
    fmt.titulo_tabla(doc, "Tabla 1.1. Datos generales del proyecto.")
    fmt.tabla(doc, ["Concepto", "Dato"], filas, tam=9, anchos_cm=[6, 10])

    doc.add_heading("1.2 Descripción del evento", level=2)
    producto = c.producto.lower() if c.producto else completar("producto derramado")
    texto = (f"El evento consistió en el derrame de {producto} en {c.f('ubicacion', 'ubicación de referencia')}, "
             f"municipio de {c.f('municipio', 'municipio')}, {c.f('estado', 'estado')}, ocurrido el "
             f"{c.f('fecha_evento', 'fecha del evento')}.")
    vol = c.mf("volumen_derramado_l")
    texto += (f" El volumen derramado reportado es de {vol:,.0f} L (fuente: {c.medio['volumen_derramado_l']['fuente']})."
              if vol is not None else " " + completar("volumen derramado y su fuente") + ".")
    acciones = c.mf("acciones_contencion")
    texto += (f" Acciones de contención: {acciones}." if acciones
              else " " + completar("acciones de contención y atención de la emergencia") + ".")
    fmt.parrafo(doc, texto)

    doc.add_heading("1.3 Objetivo", level=2)
    fmt.parrafo(doc, "Determinar la presencia y concentración de hidrocarburos en el suelo del sitio afectado, "
                     "delimitar la extensión de la afectación y comparar los resultados con los límites máximos "
                     f"permisibles (LMP) de la {NORMA} para uso de suelo {c.usos_txt()}, como base para, en su "
                     "caso, el plan de saneamiento.")


def _marco_legal(doc: Any, c: _Contexto) -> None:
    _capitulo(doc, CAPITULOS[1])
    doc.add_heading(f"2.1 {NORMA}", level=2)
    fmt.parrafo(doc, f"La {NORMA}, «{TITULO}», publicada en el Diario Oficial de la Federación el {PUBLICACION_DOF}, "
                     "establece los hidrocarburos que deben analizarse en función del producto contaminante (Tabla 1), "
                     "los límites máximos permisibles para fracciones de hidrocarburos (Tabla 2) y para hidrocarburos "
                     f"específicos (Tabla 3) en suelo, en {UNIDAD}, según el uso de suelo.")
    if c.producto:
        req = parametros_requeridos(c.producto)
        fmt.parrafo(doc, f"Para el producto involucrado ({c.producto.lower()}), la Tabla 1 de la norma establece el "
                         f"análisis de: {_lista([PARAMETROS[x].nombre.lower() for x in req])}.")
    if len(c.usos) > 1:
        fmt.parrafo(doc, "Al tratarse de un sitio con uso de suelo mixto, se aplica el límite más estricto de los usos "
                         f"involucrados ({REFERENCIAS['uso_mixto']}).")
    doc.add_heading("2.2 Otras disposiciones aplicables", level=2)
    fmt.parrafo(doc, completar("otras disposiciones jurídicas aplicables que determine el responsable técnico; la "
                               "plataforma solo incorpora la NOM-138-SEMARNAT/SSA1-2012"))


def _metodologia(doc: Any, c: _Contexto) -> None:
    _capitulo(doc, CAPITULOS[2])
    doc.add_heading("3.1 Distribución de los puntos de muestreo", level=2)
    zonas: dict[str, int] = {}
    for pt in c.puntos:
        zonas[pt["zona"] or "Sin zona"] = zonas.get(pt["zona"] or "Sin zona", 0) + 1
    muestras_por_punto: dict[str, list[dict[str, Any]]] = {}
    for f in c.filas:
        muestras_por_punto.setdefault(f["punto"] or "—", [])
        if f["muestra"] not in [m["muestra"] for m in muestras_por_punto[f["punto"] or "—"]]:
            muestras_por_punto[f["punto"] or "—"].append(f)
    n_muestras = len({f["muestra"] for f in c.filas})
    if c.puntos:
        fmt.parrafo(doc, f"Se establecieron {len(c.puntos)} puntos de muestreo, distribuidos en "
                         f"{_lista([f'{z} ({n})' for z, n in zonas.items()])}, de los que se obtuvieron "
                         f"{n_muestras} muestras de suelo (Tabla 3.1). Las coordenadas se expresan en proyección UTM"
                         + (f", zona {c.ficha['utm_zona']}" if c.ficha.get("utm_zona") else
                            ", " + completar("zona UTM y datum")) + ".")
        fmt.titulo_tabla(doc, "Tabla 3.1. Puntos de muestreo.")
        fmt.tabla(doc, ["Punto", "Zona", "X (UTM)", "Y (UTM)", "Muestras", "Profundidades (m)"], [
            [pt["clave"], pt["zona"], f"{pt['utm_x']:.2f}" if pt["utm_x"] is not None else completar("X"),
             f"{pt['utm_y']:.2f}" if pt["utm_y"] is not None else completar("Y"),
             len(muestras_por_punto.get(pt["clave"], [])),
             ", ".join(sorted({str(m["profundidad_m"]) for m in muestras_por_punto.get(pt["clave"], [])
                               if m["profundidad_m"]})) or completar("profundidad")]
            for pt in c.puntos], tam=8)
    else:
        fmt.parrafo(doc, completar("puntos de muestreo: cargue los resultados del laboratorio en el Módulo 2"))

    doc.add_heading("3.2 Toma de muestras", level=2)
    fechas = sorted({str(f["fecha_muestreo"]) for f in c.filas if f["fecha_muestreo"]})
    fmt.parrafo(doc, (f"El muestreo se realizó el {_lista(fechas)}. " if fechas
                      else completar("fechas de muestreo") + ". ")
                + completar("técnica de muestreo, equipo utilizado, preservación de las muestras y criterios de "
                            "selección de los puntos"))

    doc.add_heading("3.3 Análisis de laboratorio", level=2)
    lab = c.f("laboratorio", "laboratorio que realizó los análisis")
    acred = c.ficha.get("acreditacion_laboratorio")
    fmt.parrafo(doc, f"Las muestras se analizaron en {lab}, "
                     + (f"laboratorio acreditado ({acred})." if acred
                        else completar("acreditación vigente del laboratorio (organismo y número)") + ".")
                     + " Los parámetros analizados y sus métodos se presentan en la Tabla 3.2.")
    metodos: dict[str, set[str]] = {}
    for f in c.filas:
        if f["parametro"]:
            metodos.setdefault(f["parametro"], set())
            if f["metodo"]:
                metodos[f["parametro"]].add(f["metodo"])
    fmt.titulo_tabla(doc, "Tabla 3.2. Parámetros y métodos analíticos.")
    fmt.tabla(doc, ["Parámetro", "Método analítico", "Muestras analizadas"], [
        [PARAMETROS[k].nombre if k in PARAMETROS else k, ", ".join(sorted(v)) or completar("método analítico"),
         sum(1 for f in c.filas if f["parametro"] == k)]
        for k, v in sorted(metodos.items(), key=lambda x: list(PARAMETROS).index(x[0]) if x[0] in PARAMETROS else 99)],
        tam=8)
    faltan = [m for m in c.evaluadas if m.faltantes]
    if faltan:
        fmt.parrafo(doc, f"En {len(faltan)} muestra(s) no se reportan todos los parámetros que exige la Tabla 1 de la "
                         f"{NORMA} para el producto: "
                         + "; ".join(f"{m.muestra} ({', '.join(PARAMETROS[x].nombre.lower() for x in m.faltantes)})"
                                     for m in faltan[:15]) + ".")

    doc.add_heading("3.4 Cadena de custodia", level=2)
    fmt.parrafo(doc, completar("manejo, preservación y cadena de custodia de las muestras, y referencia al anexo con "
                               "las cadenas de custodia firmadas"))


def _resultados(doc: Any, c: _Contexto) -> None:
    _capitulo(doc, CAPITULOS[3])
    if not c.evaluadas:
        fmt.parrafo(doc, completar("resultados de laboratorio y uso de suelo del proyecto (Módulos 1 y 2)"))
        return
    doc.add_heading("4.1 Límites máximos permisibles aplicables", level=2)
    presentes = [k for k in PARAMETROS if any(f["parametro"] == k for f in c.filas)]
    fmt.parrafo(doc, f"Los resultados se comparan con los LMP de la {NORMA} para uso de suelo {c.usos_txt()} "
                     "(Tabla 4.1).")
    fmt.titulo_tabla(doc, f"Tabla 4.1. Límites máximos permisibles aplicables ({UNIDAD}).")
    fmt.tabla(doc, ["Parámetro", "LMP", "Referencia"],
              [[PARAMETROS[k].nombre, f"{lmp_aplicable(k, c.usos):g}", f"{NORMA}, {PARAMETROS[k].tabla}"]
               for k in presentes], tam=9)

    doc.add_heading("4.2 Resultados analíticos", level=2)
    fmt.parrafo(doc, "La Tabla 4.2 presenta los resultados del laboratorio por muestra y su comparación con los LMP.")
    fmt.seccion_horizontal(doc)
    escribir_tabla_nom138(doc, c.evaluadas, c.usos, c.ficha, numero_tabla="4.2")
    fmt.seccion_vertical(doc)

    doc.add_heading("4.3 Interpretación de resultados", level=2)
    s = c.resumen
    pe = s["por_estado"]
    texto = (f"De las {s['muestras']} muestras evaluadas, {pe.get(EXCEDE, 0)} exceden al menos un LMP, "
             f"{pe.get(CUMPLE, 0)} cumplen")
    if pe.get(INDETERMINADO):
        texto += (f" y {pe[INDETERMINADO]} tienen dictamen indeterminado (parámetros de la Tabla 1 no reportados o "
                  "límite de detección superior al LMP)")
    texto += "."
    if s["puntos_con_excedencia"]:
        texto += f" Los puntos con excedencia son: {_lista(s['puntos_con_excedencia'])}."
    fmt.parrafo(doc, texto)
    exced = [(m, r) for m in c.evaluadas for r in m.resultados if r.estado == EXCEDE]
    if exced:
        fmt.titulo_tabla(doc, "Tabla 4.3. Muestras con concentraciones superiores al LMP.")
        fmt.tabla(doc, ["Muestra", "Punto", "Zona", "Prof. (m)", "Parámetro", f"Concentración ({UNIDAD})", "LMP",
                        "Veces el LMP"],
                  [[m.muestra, m.punto, m.zona, m.datos.get("profundidad_m"), PARAMETROS[r.parametro].nombre,
                    formatear_valor(r), f"{r.lmp:g}", f"{r.proporcion:.1f}"] for m, r in exced], tam=8)


def _pluma(doc: Any, c: _Contexto) -> None:
    _capitulo(doc, CAPITULOS[5])
    a = c.a
    if not a.get("disponible"):
        fmt.parrafo(doc, completar("identificación de la pluma: requiere resultados de laboratorio"))
        return
    e, d = a["espacial"], a["delimitacion"]
    exced = [x for x in a["puntos"] if x["excede"]]
    if not exced:
        fmt.parrafo(doc, f"No se identificaron concentraciones superiores a los LMP de la {NORMA}; no se delimita una "
                         "pluma contaminante.")
        return
    ind = a["indicador"]["nombre"].lower() if a["indicador"] else "—"
    partes = [f"La pluma contaminante se identifica a partir de los {len(exced)} puntos con excedencia de LMP, con "
              f"la {ind} como parámetro indicador (Tabla 6.1 y Figura 5.1)."]
    if "extension_max_m" in e:
        partes.append(f"Se extiende hasta {e['extension_max_m']:g} m desde el origen"
                      + (f", con dirección predominante hacia el {e['direccion_dispersion']['rumbo']}"
                         if "direccion_dispersion" in e else "") + ".")
    if "area_afectada_m2" in e:
        partes.append(f"El área de la envolvente es de {e['area_afectada_m2']:,.0f} m²")
        if "volumen_estimado_m3" in e:
            v = e["volumen_estimado_m3"]
            partes[-1] += (f", con profundidad máxima de afectación de {e['prof_max_afectacion_m']:g} m y un volumen "
                           f"preliminar de {v['central']:,.0f} a {v['conservador']:,.0f} m³")
        partes[-1] += "."
    for eje in ("horizontal", "vertical"):
        partes.append(f"Delimitación {eje}: " + ("cerrada." if d[eje]["cerrada"] else
                                                f"abierta; {d[eje]['motivo'][0].lower()}{d[eje]['motivo'][1:]}. "
                                                "Se requieren puntos adicionales para completarla."))
    fmt.parrafo(doc, " ".join(partes))
    fmt.titulo_tabla(doc, "Tabla 6.1. Puntos que conforman la pluma contaminante.")
    fmt.tabla(doc, ["Punto", "Zona", "Parámetros que exceden", f"{ind.capitalize()} máx. ({UNIDAD})",
                    "Prof. máx. con excedencia (m)", "Delimitación vertical"],
              [[x["punto"], x["zona"], ", ".join(PARAMETROS[k].nombre for k in x["parametros_excedidos"]),
                f"{x['indicador_max']:g}" if x["indicador_max"] is not None else None, x["prof_max_excedencia_m"],
                "Abierta" if x["vertical_abierta"] else "Cerrada"] for x in exced], tam=8)


def _danos(doc: Any, c: _Contexto) -> None:
    _capitulo(doc, CAPITULOS[6])
    a = c.a
    niveles = {x["factor"]: x for x in a.get("matriz", [])}
    e = a.get("espacial", {})
    doc.add_heading("7.1 Suelo", level=2)
    if c.resumen and c.resumen["por_estado"].get(EXCEDE):
        maxima = max((r.proporcion for m in c.evaluadas for r in m.resultados if r.estado == EXCEDE), default=None)
        texto = (f"El suelo presenta concentraciones de hidrocarburos superiores a los LMP en "
                 f"{len(c.resumen['puntos_con_excedencia'])} de los {len(c.puntos)} puntos de muestreo, hasta "
                 f"{maxima:.1f} veces el LMP.")
        if "area_afectada_m2" in e:
            texto += f" Área afectada estimada: {e['area_afectada_m2']:,.0f} m²."
        if "volumen_estimado_m3" in e:
            v = e["volumen_estimado_m3"]
            texto += f" Volumen preliminar de suelo afectado: {v['central']:,.0f} a {v['conservador']:,.0f} m³."
        fmt.parrafo(doc, texto)
    else:
        fmt.parrafo(doc, "No se identificaron concentraciones superiores a los LMP en el suelo muestreado.")
    for num, factor, titulo in (("7.2", "Vulnerabilidad del agua subterránea", "Agua subterránea"),
                                ("7.3", "Afectación de cuerpos de agua superficiales", "Agua superficial")):
        doc.add_heading(f"{num} {titulo}", level=2)
        x = niveles.get(factor)
        texto = (f"Susceptibilidad {x['nivel']} ({x['justificacion']}), conforme a la matriz de la Tabla 5.3. "
                 if x else "")
        fmt.parrafo(doc, texto + completar(f"resultados o evidencias de afectación al {titulo.lower()}, en su caso"))
    doc.add_heading("7.4 Vegetación y uso del suelo", level=2)
    veg = c.mf("cobertura_vegetal")
    fmt.parrafo(doc, (f"Cobertura vegetal del sitio: {veg}. " if veg else "")
                + completar("afectación observada a la vegetación, cultivos o uso del predio"))


def _conclusiones(doc: Any, c: _Contexto) -> None:
    _capitulo(doc, CAPITULOS[7])
    concl: list[str] = []
    n_m = len({f["muestra"] for f in c.filas})
    if c.puntos:
        concl.append(f"Se caracterizaron {len(c.puntos)} puntos de muestreo con {n_m} muestras de suelo, evaluadas "
                     f"conforme a la {NORMA} para uso de suelo {c.usos_txt()}.")
    if c.resumen:
        pe = c.resumen["por_estado"]
        if pe.get(EXCEDE):
            params = sorted({x for m in c.evaluadas for x in m.excedencias}, key=list(PARAMETROS).index)
            concl.append(f"{pe[EXCEDE]} muestra(s) exceden los LMP de {_lista([PARAMETROS[x].nombre.lower() for x in params])}"
                         f", en los puntos {_lista(c.resumen['puntos_con_excedencia'])}.")
        else:
            concl.append("Ninguna muestra presenta concentraciones superiores a los LMP aplicables.")
        if pe.get(INDETERMINADO):
            concl.append(f"{pe[INDETERMINADO]} muestra(s) tienen dictamen indeterminado y requieren completar o "
                         "aclarar sus resultados analíticos.")
    d = c.a.get("delimitacion")
    if d and c.resumen and c.resumen["por_estado"].get(EXCEDE):
        abiertas = [eje for eje in ("horizontal", "vertical") if not d[eje]["cerrada"]]
        concl.append("La pluma contaminante se encuentra delimitada." if not abiertas else
                     f"La delimitación {' y '.join(abiertas)} de la pluma no está cerrada; se requiere muestreo "
                     "complementario.")
    altas = [x["factor"].lower() for x in c.a.get("matriz", []) if x["nivel"] == "alta"]
    if altas:
        concl.append(f"Los factores con susceptibilidad alta a la migración son: {_lista(altas)}.")
    if c.resumen and c.resumen["por_estado"].get(EXCEDE):
        concl.append(f"Se requiere el saneamiento del suelo afectado hasta alcanzar los LMP de la {NORMA} para el "
                     "uso de suelo del sitio (Capítulo 9).")
    if not concl:
        concl.append(completar("conclusiones"))
    for i, t in enumerate(concl, start=1):
        fmt.parrafo(doc, f"{i}. {t}")


def _saneamiento(doc: Any, c: _Contexto) -> None:
    _capitulo(doc, CAPITULOS[8])
    exced = c.resumen and c.resumen["por_estado"].get(EXCEDE)
    if not exced:
        fmt.parrafo(doc, "Con base en los resultados no se identifican concentraciones superiores a los LMP; no se "
                         "propone saneamiento. " + completar("confirmación del responsable técnico"))
        return
    doc.add_heading("9.1 Objetivo y metas de limpieza", level=2)
    fmt.parrafo(doc, f"Reducir las concentraciones de hidrocarburos en el suelo afectado por debajo de los LMP de la "
                     f"{NORMA} para uso de suelo {c.usos_txt()}. Las metas de limpieza se presentan en la Tabla 9.1.")
    params = sorted({x for m in c.evaluadas for x in m.excedencias}, key=list(PARAMETROS).index)
    fmt.titulo_tabla(doc, f"Tabla 9.1. Metas de limpieza ({UNIDAD}).")
    fmt.tabla(doc, ["Parámetro", "Concentración máxima registrada", "Meta (LMP)", "Referencia"], [
        [PARAMETROS[k].nombre, f"{c.resumen['maximos'][k]:g}", f"{lmp_aplicable(k, c.usos):g}",
         f"{NORMA}, {PARAMETROS[k].tabla}"] for k in params], tam=9)
    doc.add_heading("9.2 Área y volumen a sanear", level=2)
    e = c.a.get("espacial", {})
    if "volumen_estimado_m3" in e:
        v = e["volumen_estimado_m3"]
        texto = (f"Con base en la envolvente de puntos con excedencia ({e['area_afectada_m2']:,.0f} m²), el volumen "
                 f"preliminar de suelo a sanear es de {v['central']:,.0f} a {v['conservador']:,.0f} m³.")
        d = c.a["delimitacion"]
        if not (d["horizontal"]["cerrada"] and d["vertical"]["cerrada"]):
            texto += " La estimación es preliminar: la delimitación de la pluma no está cerrada."
        fmt.parrafo(doc, texto)
    else:
        fmt.parrafo(doc, completar("área y volumen de suelo a sanear"))
    for num, titulo, falta in (("9.3", "Técnica de remediación", "técnica de remediación seleccionada y su justificación"),
                               ("9.4", "Muestreo de comprobación", "plan de muestreo para comprobar el cumplimiento de "
                                                                  "las metas de limpieza"),
                               ("9.5", "Programa de actividades", "programa calendarizado de actividades")):
        doc.add_heading(f"{num} {titulo}", level=2)
        fmt.parrafo(doc, completar(falta))


def _anexos(doc: Any, c: _Contexto) -> None:
    _capitulo(doc, CAPITULOS[10])
    labs = [d["nombre"] for d in c.p.documentos("laboratorio")]
    for etq, contenido in (
        ("Anexo A. Informe de resultados del laboratorio", _lista(labs) if labs else completar("informe oficial del laboratorio")),
        ("Anexo B. Cadenas de custodia", completar("cadenas de custodia firmadas")),
        ("Anexo C. Planos", completar("plano de localización del sitio y plano de puntos de muestreo")),
    ):
        fmt.parrafo(doc, f"{etq}: {contenido}.")


# ------------------------------------------------------------------ verificación previa
def verificacion_previa(p: Proyecto) -> list[dict[str, str]]:
    """Estado de los insumos del informe: «ok», «aviso» (se generará con pendientes) o «falta»."""
    c = _Contexto(p)
    v: list[dict[str, str]] = []

    def add(nivel: str, texto: str) -> None:
        v.append({"nivel": nivel, "texto": texto})

    falt = campos_faltantes(c.ficha)
    add("aviso" if falt else "ok", "Ficha del proyecto completa" if not falt else
        "Ficha incompleta: " + ", ".join(x.etiqueta for x in falt))
    if not c.evaluadas:
        add("falta", "Sin resultados de laboratorio evaluados (Módulos 1–3)")
    else:
        esperadas = c.ficha.get("muestras_esperadas")
        n = len(c.evaluadas)
        add("aviso" if esperadas and str(n) != str(esperadas).strip() else "ok",
            f"{n} muestras evaluadas" + (f" (esperadas: {esperadas})" if esperadas else ""))
        faltan = sum(1 for m in c.evaluadas if m.faltantes)
        if faltan:
            add("aviso", f"{faltan} muestra(s) sin todos los parámetros de la Tabla 1")
    fotos = diag_fotos(p)
    add("aviso" if not fotos["total"] or fotos["sin_clasificar"] or fotos["sin_revisar"] else "ok",
        f"Fotografías: {fotos['total']} registradas, {len(fotos['sin_clasificar'])} sin clasificar, "
        f"{len(fotos['sin_revisar'])} sin revisar")
    req = [x for x in mf.CAMPOS if x.requerido]
    sin = [x.etiqueta for x in req if c.mf(x.clave) is None]
    add("aviso" if sin else "ok", "Medio físico completo con fuentes" if not sin else
        f"Medio físico: {len(sin)} dato(s) requerido(s) sin valor o sin fuente")
    if c.cap5_guardado:
        n = len(c.cap5["cifras_no_respaldadas"])
        add("aviso" if n else "ok", f"Capítulo 5 versión #{c.cap5['id']}"
            + (f" con {n} cifra(s) sin respaldo resaltadas" if n else ""))
    else:
        add("aviso", "Capítulo 5 no generado: se incluirá la redacción por plantilla")
    d = c.a.get("delimitacion") if c.a.get("disponible") else None
    if d and c.resumen and c.resumen["por_estado"].get(EXCEDE):
        for eje in ("horizontal", "vertical"):
            if not d[eje]["cerrada"]:
                add("aviso", f"Delimitación {eje} abierta")
    revs = listar_revisiones(p)
    if revs:
        criticos = sum(1 for h in hallazgos_de(p, revs[0]["id"]) if h.estado == "abierto" and h.criticidad == "critica")
        add("aviso" if criticos else "ok", f"Última revisión técnica: ICT {revs[0]['ict_vigente']:.0f}"
            + (f", {criticos} hallazgo(s) crítico(s) abiertos" if criticos else ""))
    return v


# ------------------------------------------------------------------ generación
def generar(p: Proyecto, plantilla: Path | bytes | None = None, hoy: date | None = None) -> tuple[bytes, dict[str, Any]]:
    """Genera el informe. Devuelve (docx, metadatos con pendientes y avisos)."""
    hoy = hoy or date.today()
    c = _Contexto(p)
    doc = fmt.documento(plantilla)
    _portada(doc, c, hoy)
    fmt.salto_pagina(doc)
    fmt.indice(doc)
    _antecedentes(doc, c)
    _marco_legal(doc, c)
    _metodologia(doc, c)
    _resultados(doc, c)
    fmt.salto_pagina(doc)
    cap5.escribir_capitulo(doc, c.cap5)
    _pluma(doc, c)
    _danos(doc, c)
    _conclusiones(doc, c)
    _saneamiento(doc, c)
    fmt.salto_pagina(doc)
    avisos_fotos, n_fotos = escribir_seccion(doc, p, "Memoria fotográfica", capitulo="10")
    _anexos(doc, c)
    fmt.encabezado(doc, f"{c.ficha.get('nombre') or p.clave} · Informe de Caracterización")
    fmt.pie_paginado(doc, CONSULTORA)
    fmt.actualizar_campos_al_abrir(doc)
    fmt.precargar_indice(doc)
    pendientes = fmt.contar_pendientes(doc)
    buf = io.BytesIO()
    doc.save(buf)
    meta = {
        "pendientes": pendientes, "fotografias": n_fotos, "avisos": avisos_fotos
        + ([f"Capítulo 5: {len(c.cap5['cifras_no_respaldadas'])} cifra(s) sin respaldo resaltadas"]
           if c.cap5["cifras_no_respaldadas"] else [])
        + ([] if c.cap5_guardado else ["Capítulo 5 por plantilla (no se ha generado una versión en el Módulo 6)"]),
    }
    return buf.getvalue(), meta


def generar_y_registrar(p: Proyecto, plantilla_id: int | None = None, revisar: bool = True,
                        hoy: date | None = None) -> dict[str, Any]:
    """Genera el informe, lo guarda como borrador del proyecto y, opcionalmente, lo somete a la revisión técnica."""
    from ..revision.registro import revisar_borrador

    plantilla = None
    if plantilla_id is not None:
        doc_p = next((d for d in p.documentos() if d["id"] == plantilla_id), None)
        if doc_p is None:
            raise KeyError(f"No existe la plantilla {plantilla_id}")
        plantilla = Path(doc_p["ruta"])
    contenido, meta = generar(p, plantilla, hoy)
    previos = [d for d in p.documentos("borrador") if d["nombre"].startswith("Informe_caracterizacion_")]
    nombre = f"Informe_caracterizacion_{p.clave}_v{len(previos) + 1}.docx"
    p.guardar_documento("borrador", nombre, contenido)
    doc_id = p.documentos("borrador")[0]["id"]
    p.registrar_cambio("informe", nombre, "generar", None,
                       {"documento_id": doc_id, "pendientes": len(meta["pendientes"])})
    p.con.commit()
    out = {"documento_id": doc_id, "nombre": nombre, "ruta": next(d["ruta"] for d in p.documentos()
                                                                  if d["id"] == doc_id), **meta}
    if revisar:
        rev_id = revisar_borrador(p, doc_id, usar_ia=False, hoy=hoy)
        rev = next(r for r in listar_revisiones(p) if r["id"] == rev_id)
        out.update({"revision_id": rev_id, "ict": rev["ict"], "interpretacion": rev["interpretacion"]})
    return out
