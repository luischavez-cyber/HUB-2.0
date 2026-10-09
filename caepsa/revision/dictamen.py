"""Exportación del dictamen de revisión técnica (Word y Excel)."""
from __future__ import annotations

import io
from typing import Any

from ..proyectos import Proyecto
from .indice import CATEGORIAS, CRITICIDADES, DIMENSIONES, ETIQUETA_CRITICIDAD, Hallazgo, calcular_ict
from .registro import hallazgos_de, obtener_revision

COLOR = {"critica": "F8D7DA", "alta": "FCE4D6", "media": "FFF3CD", "baja": "E2EFDA", "info": "EDEDED"}
AVISO = ("Índice de Calidad Técnica (ICT): indicador interno de CAEPSA calculado a partir de los hallazgos. "
         "No sustituye la evaluación ni la resolución de la autoridad.")


def _ordenados(hallazgos: list[Hallazgo]) -> list[Hallazgo]:
    orden = list(CRITICIDADES)
    return sorted([h for h in hallazgos if h.estado != "descartado"], key=lambda h: (orden.index(h.criticidad),
                                                                                     h.origen, h.id or 0))


def _evidencia(h: Hallazgo) -> str:
    if not h.evidencia:
        return ""
    sufijo = " (cita no localizada en el documento; verificar)" if h.evidencia_localizada is False else ""
    return f"«{h.evidencia}»{sufijo}"


def a_word(p: Proyecto, revision_id: int) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    rev = obtener_revision(p, revision_id)
    hallazgos = hallazgos_de(p, revision_id)
    ict = calcular_ict(hallazgos)
    ficha = p.ficha()

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for lado in ("left_margin", "right_margin"):
        setattr(sec, lado, Cm(1.8))
    doc.add_heading("Dictamen de revisión técnica del informe", level=1)
    for etq, val in (("Proyecto", ficha.get("nombre")), ("Expediente", ficha.get("expediente")),
                     ("Documento revisado", rev["documento"]), ("Fecha de revisión", rev["fecha"]),
                     ("Revisó", rev["usuario"]), ("Revisión con IA", "Sí" if rev["usa_ia"] else "No")):
        par = doc.add_paragraph()
        par.add_run(f"{etq}: ").bold = True
        par.add_run(str(val or "—"))
    if rev["aviso"]:
        doc.add_paragraph(rev["aviso"]).runs[0].italic = True

    doc.add_heading(f"Índice de Calidad Técnica: {ict['ict']:.0f} / 100 — {ict['interpretacion']}", level=2)
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for c, txt in zip(t.rows[0].cells, ("Dimensión", "Ponderación", "Subíndice")):
        c.text = txt
        c.paragraphs[0].runs[0].bold = True
    for d, (nombre, peso) in DIMENSIONES.items():
        fila = t.add_row().cells
        fila[0].text, fila[1].text, fila[2].text = nombre, f"{peso * 100:.0f} %", f"{ict['subindices'][d]:.0f}"
    doc.add_paragraph(AVISO).runs[0].font.size = Pt(8)

    if rev["resumen"]:
        doc.add_heading("Síntesis de la revisión", level=2)
        doc.add_paragraph(rev["resumen"])
    if rev["fortalezas"]:
        doc.add_heading("Fortalezas", level=2)
        for f in rev["fortalezas"]:
            doc.add_paragraph(f, style="List Bullet")

    lista = _ordenados(hallazgos)
    doc.add_heading(f"Hallazgos ({len(lista)})", level=2)
    enc = ("#", "Criticidad", "Categoría", "Ubicación", "Hallazgo y evidencia", "Riesgo regulatorio",
           "Recomendación", "Referencia", "Estado")
    t = doc.add_table(rows=1, cols=len(enc))
    t.style = "Table Grid"
    for c, txt in zip(t.rows[0].cells, enc):
        c.text = txt
        c.paragraphs[0].runs[0].bold = True
    for i, h in enumerate(lista, start=1):
        valores = (str(i), ETIQUETA_CRITICIDAD[h.criticidad], CATEGORIAS[h.categoria][0], h.ubicacion,
                   h.descripcion + (f"\n{_evidencia(h)}" if h.evidencia else ""), h.riesgo, h.recomendacion,
                   h.referencia, h.estado + (f": {h.nota}" if h.nota else ""))
        celdas = t.add_row().cells
        for c, txt in zip(celdas, valores):
            c.text = txt
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), COLOR[h.criticidad])
        celdas[1]._tc.get_or_add_tcPr().append(shd)
    for fila in t.rows:
        for c in fila.cells:
            for par in c.paragraphs:
                for r in par.runs:
                    r.font.size = Pt(8)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def a_excel(p: Proyecto, revision_id: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    rev = obtener_revision(p, revision_id)
    hallazgos = hallazgos_de(p, revision_id)
    ict = calcular_ict(hallazgos)
    wb = Workbook()
    ws = wb.active
    ws.title = "Hallazgos"
    enc = ["ID", "Origen", "Criticidad", "Categoría", "Dimensión", "Ubicación", "Descripción", "Evidencia",
           "Evidencia localizada", "Riesgo regulatorio", "Recomendación", "Referencia", "Estado", "Nota"]
    ws.append(enc)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for h in sorted(hallazgos, key=lambda h: (list(CRITICIDADES).index(h.criticidad), h.id or 0)):
        ws.append([h.id, "IA" if h.origen == "ia" else "Automática", ETIQUETA_CRITICIDAD[h.criticidad],
                   CATEGORIAS[h.categoria][0], DIMENSIONES[h.dimension][0], h.ubicacion, h.descripcion, h.evidencia,
                   {True: "Sí", False: "No", None: "—"}[h.evidencia_localizada], h.riesgo, h.recomendacion,
                   h.referencia, h.estado, h.nota])
        ws.cell(row=ws.max_row, column=3).fill = PatternFill("solid", fgColor=COLOR[h.criticidad])
    for col, ancho in zip("ABCDEFGHIJKLMN", (6, 11, 11, 26, 22, 22, 60, 50, 11, 40, 50, 28, 11, 30)):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    res = wb.create_sheet("Índice")
    res.append(["Documento", rev["documento"]])
    res.append(["Fecha", rev["fecha"]])
    res.append(["ICT original", rev["ict"]])
    res.append(["ICT vigente (sin descartados)", ict["ict"]])
    res.append(["Interpretación", ict["interpretacion"]])
    for d, (nombre, peso) in DIMENSIONES.items():
        res.append([f"{nombre} ({peso * 100:.0f} %)", ict["subindices"][d]])
    res.append([AVISO])
    res.column_dimensions["A"].width = 40
    res.column_dimensions["B"].width = 50
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
