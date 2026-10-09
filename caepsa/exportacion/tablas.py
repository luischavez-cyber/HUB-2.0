"""Exportación de la evaluación NOM-138 a Excel y Word."""
from __future__ import annotations

import io
from typing import Any

from ..normativa.evaluacion import EXCEDE, INDETERMINADO, MuestraEvaluada, resumen
from ..normativa.nom138 import NORMA, PARAMETROS, PUBLICACION_DOF, REFERENCIAS, UNIDAD, UsoSuelo, lmp_aplicable

ROJO, AMBAR, VERDE = "F8D7DA", "FFF3CD", "D4EDDA"


def formatear_valor(r: Any) -> str:
    if r is None:
        return "—"
    if r.menor_que:
        umbral = r.valor if r.valor is not None else r.limite_deteccion
        return f"<{umbral:g}" if umbral is not None else "ND"
    return "—" if r.valor is None else f"{r.valor:,.2f}".rstrip("0").rstrip(".")


def parametros_presentes(muestras: list[MuestraEvaluada]) -> list[str]:
    presentes = {r.parametro for m in muestras for r in m.resultados if r.parametro in PARAMETROS}
    return [c for c in PARAMETROS if c in presentes]


def matriz(muestras: list[MuestraEvaluada]) -> tuple[list[str], list[list[Any]], list[list[str | None]]]:
    """Encabezados, valores y estados (para colorear) — una fila por muestra."""
    params = parametros_presentes(muestras)
    enc = ["Zona", "Punto", "Muestra", "Prof. (m)", "X (UTM)", "Y (UTM)",
           *[PARAMETROS[c].nombre for c in params], "Dictamen"]
    filas, estados = [], []
    for m in muestras:
        por_p = {r.parametro: r for r in m.resultados}
        d = m.datos
        filas.append([m.zona or "", m.punto or "", m.muestra, d.get("profundidad_m") or "",
                      d.get("utm_x") or "", d.get("utm_y") or "",
                      *[formatear_valor(por_p.get(c)) for c in params], m.estado])
        estados.append([None] * 6 + [por_p[c].estado if c in por_p else None for c in params] + [m.estado])
    return enc, filas, estados


def _fila_lmp(params: list[str], usos: list[UsoSuelo]) -> list[Any]:
    return ["", "", "LMP NOM-138", "", "", "", *[lmp_aplicable(c, usos) for c in params], ""]


def _color(estado: str | None) -> str | None:
    return {EXCEDE: ROJO, INDETERMINADO: AMBAR}.get(estado or "")


def a_excel(muestras: list[MuestraEvaluada], usos: list[UsoSuelo], ficha: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Cumplimiento NOM-138"
    enc, filas, estados = matriz(muestras)
    params = parametros_presentes(muestras)
    borde = Border(*(Side(style="thin", color="999999"),) * 4)

    ws.append([f"Evaluación de cumplimiento — {NORMA} (DOF {PUBLICACION_DOF})"])
    ws.append([f"Proyecto: {ficha.get('nombre', '')}   Expediente: {ficha.get('expediente', '')}"])
    ws.append([f"Uso de suelo: {' / '.join(u.etiqueta for u in usos)}   Unidades: {UNIDAD}"])
    ws.append([])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append(enc)
    ws.append(_fila_lmp(params, usos))
    for f in filas:
        ws.append(f)
    inicio = 5
    for c in ws[inicio]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for c in ws[inicio + 1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for i, est in enumerate(estados):
        for j, e in enumerate(est):
            celda = ws.cell(row=inicio + 2 + i, column=j + 1)
            color = _color(e) if j < len(est) - 1 else {EXCEDE: ROJO, INDETERMINADO: AMBAR}.get(e or "", VERDE)
            if color:
                celda.fill = PatternFill("solid", fgColor=color)
            if e == EXCEDE:
                celda.font = Font(bold=True, color="9C0006")
    for fila in ws.iter_rows(min_row=inicio, max_row=ws.max_row):
        for c in fila:
            c.border = borde
    for j in range(1, len(enc) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 14
    ws.freeze_panes = ws.cell(row=inicio + 2, column=4)

    # Hoja de detalle con trazabilidad completa.
    det = wb.create_sheet("Detalle")
    det.append(["Muestra", "Punto", "Zona", "Parámetro", "Valor reportado", "LD", "Unidad", "Método", "LMP",
                "Valor/LMP", "Estado", "Referencia normativa", "Origen", "Página PDF", "Observaciones"])
    for m in muestras:
        for r in m.resultados:
            det.append([m.muestra, m.punto, m.zona, PARAMETROS[r.parametro].nombre if r.parametro in PARAMETROS
                        else r.parametro, formatear_valor(r), r.limite_deteccion, r.datos.get("unidad"),
                        r.datos.get("metodo"), r.lmp, round(r.proporcion, 3) if r.proporcion is not None else None,
                        r.estado, r.referencia, r.datos.get("origen"), r.datos.get("pagina_pdf"),
                        "; ".join(r.observaciones)])
            if _color(r.estado):
                det.cell(row=det.max_row, column=11).fill = PatternFill("solid", fgColor=_color(r.estado))
        if m.faltantes:
            det.append([m.muestra, m.punto, m.zona, "Parámetros faltantes (Tabla 1)", ", ".join(m.faltantes)])

    res = wb.create_sheet("Resumen")
    s = resumen(muestras)
    for k, v in [("Muestras evaluadas", s["muestras"]), ("Puntos de muestreo", s["puntos"]),
                 *[(f"Muestras — {e}", n) for e, n in s["por_estado"].items()],
                 ("Puntos con excedencia", ", ".join(s["puntos_con_excedencia"])),
                 *[(f"Excedencias {PARAMETROS[p].nombre}", n) for p, n in s["excedencias_por_parametro"].items()],
                 ("Referencias", "; ".join(REFERENCIAS.values()))]:
        res.append([k, v])
    res.column_dimensions["A"].width = 34
    res.column_dimensions["B"].width = 80

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def a_word(muestras: list[MuestraEvaluada], usos: list[UsoSuelo], ficha: dict[str, Any],
           numero_tabla: str = "4.1") -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    for lado in ("left_margin", "right_margin"):
        setattr(sec, lado, Cm(1.5))

    enc, filas, estados = matriz(muestras)
    params = parametros_presentes(muestras)
    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titulo.add_run(f"Tabla {numero_tabla}. Resultados analíticos de suelo y comparación con los límites "
                         f"máximos permisibles de la {NORMA} ({UNIDAD}).")
    run.bold = True
    run.font.size = Pt(10)

    tabla = doc.add_table(rows=1, cols=len(enc))
    tabla.style = "Table Grid"

    def sombrear(celda: Any, color: str) -> None:
        tc = celda._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), color)
        tc.append(shd)

    def escribir(celda: Any, texto: Any, negrita: bool = False, rojo: bool = False) -> None:
        celda.text = ""
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("" if texto is None else str(texto))
        r.font.size = Pt(7)
        r.bold = negrita
        if rojo:
            r.font.color.rgb = RGBColor(0x9C, 0x00, 0x06)

    for j, h in enumerate(enc):
        escribir(tabla.rows[0].cells[j], h, negrita=True)
        sombrear(tabla.rows[0].cells[j], "D9D9D9")
    lmp = tabla.add_row().cells
    for j, v in enumerate(_fila_lmp(params, usos)):
        escribir(lmp[j], f"{v:g}" if isinstance(v, (int, float)) else v, negrita=True)
        sombrear(lmp[j], "DDEBF7")
    for f, est in zip(filas, estados):
        celdas = tabla.add_row().cells
        for j, (v, e) in enumerate(zip(f, est)):
            escribir(celdas[j], v, negrita=e == EXCEDE, rojo=e == EXCEDE)
            if _color(e):
                sombrear(celdas[j], _color(e) or "")

    usos_txt = " / ".join(u.etiqueta for u in usos)
    notas = [
        f"LMP aplicables a uso de suelo {usos_txt}: {REFERENCIAS['lmp_fracciones']} y {REFERENCIAS['lmp_especificos']}.",
        "Celdas sombreadas en rojo y en negritas: concentración superior al LMP.",
        "<: concentración inferior al límite de detección/cuantificación reportado por el laboratorio.",
    ]
    if len(set(usos)) > 1:
        notas.append(f"Uso de suelo mixto: se aplica el LMP más estricto ({REFERENCIAS['uso_mixto']}).")
    if ficha.get("laboratorio"):
        notas.append(f"Fuente: informe de resultados de {ficha['laboratorio']}.")
    for n in notas:
        p = doc.add_paragraph(n)
        p.runs[0].font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
