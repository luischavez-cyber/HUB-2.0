"""Generación de la sección fotográfica del informe (Word)."""
from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from ..proyectos import Proyecto
from .clasificador import CATEGORIAS
from .metadatos import preparar_imagen
from .registro import listar_fotografias


def fotos_para_informe(p: Proyecto) -> tuple[list[dict[str, Any]], list[str]]:
    """Fotografías incluidas, clasificadas y con pie, en orden de informe; más avisos de las omitidas."""
    avisos, listas = [], []
    for f in listar_fotografias(p, solo_incluidas=True):
        if not f["categoria"] or not f["pie"]:
            avisos.append(f"{f['nombre']}: sin categoría o pie; se omite")
        else:
            listas.append(f)
    sin_revisar = sum(1 for f in listas if not f["revisada"])
    if sin_revisar:
        avisos.append(f"{sin_revisar} fotografía(s) incluidas sin revisión humana")
    return listas, avisos


def escribir_seccion(doc: Any, p: Proyecto, titulo: str = "Memoria fotográfica", columnas: int = 2,
                     capitulo: str | None = None) -> tuple[list[str], int]:
    """Escribe la sección fotográfica en un documento existente. Devuelve (avisos, fotografías incluidas).

    Con ``capitulo`` (p. ej. "10") los subtítulos se numeran 10.1, 10.2…, como en el informe completo.
    """
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    fotos, avisos = fotos_para_informe(p)
    doc.add_heading(f"{capitulo}. {titulo.upper()}" if capitulo else titulo, level=1)
    numero, sub = 0, 0
    ancho = Cm(16.5 / columnas - 0.5)
    for clave, (nombre_cat, _) in CATEGORIAS.items():
        grupo = [f for f in fotos if f["categoria"] == clave]
        if not grupo:
            continue
        sub += 1
        doc.add_heading(f"{capitulo}.{sub} {nombre_cat}" if capitulo else nombre_cat, level=2)
        tabla = doc.add_table(rows=0, cols=columnas)
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i in range(0, len(grupo), columnas):
            celdas = tabla.add_row().cells
            for j, f in enumerate(grupo[i:i + columnas]):
                numero += 1
                celda = celdas[j]
                p_img = celda.paragraphs[0]
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.add_run().add_picture(io.BytesIO(preparar_imagen(Path(f["ruta"]).read_bytes(), 1400)),
                                            width=ancho)
                pie = celda.add_paragraph()
                pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r1 = pie.add_run(f"Fotografía {numero}. ")
                r1.bold = True
                r1.font.size = Pt(9)
                pie.add_run(f["pie"]).font.size = Pt(9)
        doc.add_paragraph()
    if not numero:
        doc.add_paragraph("No hay fotografías clasificadas para incluir.")
    return avisos, numero


def generar_word(p: Proyecto, titulo: str = "Memoria fotográfica", columnas: int = 2) -> tuple[bytes, list[str]]:
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    for lado in ("left_margin", "right_margin"):
        setattr(doc.sections[0], lado, Cm(2))
    avisos, _ = escribir_seccion(doc, p, titulo, columnas)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), avisos
