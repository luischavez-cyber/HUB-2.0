"""Módulo 6 — Generación del Capítulo 5 (Análisis de dispersión de hidrocarburos).

Dos redacciones, ambas a partir de los mismos hechos (``analisis.analizar``):
* plantilla determinista (siempre disponible, sin IA);
* redacción técnica con Claude, con verificación de que cada cifra del texto exista en los datos.
"""
from __future__ import annotations

import io
import json
import os
import re
from typing import Any

import anthropic

from ..normativa.nom138 import NORMA, PUBLICACION_DOF
from ..proyectos import Proyecto
from . import medio_fisico as mf
from .analisis import analizar
from .figura import plano

MODELO = os.environ.get("CAEPSA_MODELO", "claude-opus-5-5")
SECCIONES = {
    "5.1": "Introducción y alcance",
    "5.2": "Caracterización del medio físico",
    "5.3": "Distribución espacial y vertical de la afectación",
    "5.4": "Mecanismos y rutas probables de migración",
    "5.5": "Escenarios de dispersión",
    "5.6": "Factores de riesgo ecológico",
    "5.7": "Limitaciones e información faltante",
    "5.8": "Conclusiones técnicas",
}
PENDIENTE = "[DATO PENDIENTE: {}]"


def _n(x: float | None, dec: int = 1) -> str:
    if x is None:
        return "—"
    return f"{x:,.{dec}f}".rstrip("0").rstrip(".") if isinstance(x, float) and not x.is_integer() else f"{x:,.0f}"


def _fmt_valor(clave: str, v: Any) -> str:
    if clave in ("origen_x", "origen_y"):
        return f"{v:.2f}"  # coordenadas UTM sin separador de miles
    if isinstance(v, float):
        return f"{v:g}" if abs(v) < 1e-2 else _n(v, 3)
    return str(v)


def _con_unidad(clave: str, d: dict[str, Any]) -> str:
    u = d["unidad"]
    return _fmt_valor(clave, d["valor"]) + (f" {u}" if u and u != "fracción" else "")


def _dato(a: dict[str, Any], clave: str) -> str | None:
    d = a["medio_fisico"].get(clave)
    return f"{_con_unidad(clave, d)} (fuente: {d['fuente']})" if d else None


def _v(a: dict[str, Any], clave: str) -> str:
    """Valor con unidad, sin repetir la fuente (ya citada en 5.2)."""
    d = a["medio_fisico"].get(clave)
    return _con_unidad(clave, d) if d else PENDIENTE.format(mf.POR_CLAVE[clave].etiqueta.lower())


def _o(a: dict[str, Any], clave: str) -> str:
    return _dato(a, clave) or PENDIENTE.format(mf.POR_CLAVE[clave].etiqueta.lower())


# ------------------------------------------------------------------ plantilla
def redactar_plantilla(a: dict[str, Any]) -> list[dict[str, str]]:
    if not a.get("disponible"):
        return [{"numero": n, "titulo": t, "texto": PENDIENTE.format(a.get("motivo", "análisis no disponible"))}
                for n, t in SECCIONES.items()]
    e, d, mfis = a["espacial"], a["delimitacion"], a["medio_fisico"]
    ind = a["indicador"]["nombre"] if a["indicador"] else "—"
    n_m = sum(x["muestras"] for x in a["puntos"])
    s: dict[str, str] = {}

    producto = f" ({a['producto'].lower()})" if a["producto"] else ""
    s["5.1"] = (
        f"En este capítulo se analiza el comportamiento en el suelo del producto derramado{producto}, con base en "
        f"los resultados analíticos de {n_m} muestras obtenidas en {len(a['puntos'])} puntos de muestreo, evaluados "
        f"conforme a la {NORMA} (DOF {PUBLICACION_DOF}) para uso de suelo {' / '.join(a['uso_suelo'])}, y en los "
        f"datos del medio físico documentados (Tabla 5.1). Volumen derramado reportado: {_o(a, 'volumen_derramado_l')}.")

    parrafos = []
    for grupo in mf.GRUPOS:
        if grupo == "Evento":
            continue
        claves = [k for k in mf.POR_CLAVE if mf.POR_CLAVE[k].grupo == grupo]
        partes = [f"{mf.POR_CLAVE[k].etiqueta.lower()}: {_con_unidad(k, mfis[k])}" for k in claves if k in mfis]
        partes += [PENDIENTE.format(mf.POR_CLAVE[k].etiqueta.lower()) for k in claves
                   if mf.POR_CLAVE[k].requerido and k not in mfis]
        if partes:
            fuentes = list(dict.fromkeys(mfis[k]["fuente"] for k in claves if k in mfis))
            texto = f"{grupo}: " + "; ".join(partes) + "."
            if fuentes:
                texto += f" Fuente{'s' if len(fuentes) > 1 else ''}: {'; '.join(fuentes)}."
            parrafos.append(texto)
    s["5.2"] = "\n\n".join(parrafos) or PENDIENTE.format("datos del medio físico")

    p3 = [f"De los {e['puntos_evaluados']} puntos evaluados, {e['puntos_con_excedencia']} presentan al menos una "
          f"concentración superior al LMP de la {NORMA}; el parámetro indicador de la afectación es la {ind.lower()} "
          "(Figura 5.1 y Tabla 5.2)."]
    if a["origen"]:
        p3.append("El origen se ubica en " + (f"las coordenadas documentadas X {a['origen']['x']:.2f}, "
                                              f"Y {a['origen']['y']:.2f}" if a["origen"]["tipo"] == "documentado"
                                              else f"el punto de mayor concentración ({a['origen'].get('punto')}), "
                                                   "al no contar con coordenadas documentadas del derrame") + ".")
    if "extension_max_m" in e:
        p3.append(f"La excedencia más alejada del origen se encuentra a {_n(e['extension_max_m'])} m.")
    if "direccion_dispersion" in e:
        txt = f"El centro de masa de la afectación se desplaza hacia el {e['direccion_dispersion']['rumbo']}"
        c = e.get("coincidencia_pendiente")
        if c:
            txt += (f", lo que {'coincide' if c['valoracion'] == 'coincide' else 'coincide parcialmente' if c['valoracion'] == 'parcial' else 'no coincide'}"
                    f" con la dirección de la pendiente ({c['direccion_pendiente']}; diferencia de "
                    f"{c['diferencia_grados']}°)")
        p3.append(txt + ".")
    if "area_afectada_m2" in e:
        p3.append(f"El área de la envolvente de los puntos con excedencia es de {_n(e['area_afectada_m2'])} m².")
    if "volumen_estimado_m3" in e:
        vol = e["volumen_estimado_m3"]
        p3.append(f"La profundidad máxima con excedencia es de {_n(e['prof_max_afectacion_m'], 2)} m y la media de "
                  f"{_n(e['prof_media_afectacion_m'], 2)} m, con un volumen preliminar de suelo afectado de "
                  f"{_n(vol['central'])} m³ (estimación central) a {_n(vol['conservador'])} m³ (conservadora), "
                  f"calculado como {vol['metodo']}.")
    p3.append("Delimitación horizontal: " + ("cerrada." if d["horizontal"]["cerrada"]
                                            else f"abierta. {d['horizontal']['motivo']}."))
    p3.append("Delimitación vertical: " + ("cerrada." if d["vertical"]["cerrada"]
                                          else f"abierta. {d['vertical']['motivo']}."))
    s["5.3"] = " ".join(p3)

    p4 = [f"{x['factor']}: susceptibilidad {x['nivel']}; {x['justificacion']}." for x in a["matriz"]]
    t = a.get("transito_vertical")
    if t:
        p4.append(f"Con la conductividad hidráulica y la porosidad efectiva documentadas, el tiempo de tránsito "
                  f"vertical hasta el nivel freático a través de {_n(t['espesor_no_saturado_m'], 2)} m de suelo no "
                  f"saturado se estima en {_n(t['dias'])} días. {t['supuesto']}")
    p4.append(f"Matriz de susceptibilidad en la Tabla 5.3. {a['criterio_tamizaje']}")
    s["5.4"] = "\n\n".join(p4)

    niveles = {x["factor"]: x["nivel"] for x in a["matriz"]}
    s["5.5"] = "\n\n".join([
        "Escenario 1, condición actual: la afectación se describe en la sección 5.3; "
        + ("la dispersión observada sigue la pendiente del terreno." if (e.get("coincidencia_pendiente") or {})
           .get("valoracion") == "coincide" else "la relación con la pendiente debe confirmarse en campo."),
        f"Escenario 2, temporada de lluvias: temporada {_v(a, 'meses_lluvia')} y precipitación media anual de "
        f"{_v(a, 'precipitacion_anual_mm')}; el arrastre por escurrimiento superficial tiene susceptibilidad "
        f"{niveles.get('Arrastre por escurrimiento superficial', '—')} y la lixiviación vertical "
        f"{niveles.get('Infiltración y lixiviación vertical', '—')}.",
        "Escenario 3, sin acciones de saneamiento: con movilidad del contaminante "
        f"{niveles.get('Movilidad del contaminante', '—')} y vulnerabilidad del agua subterránea "
        f"{niveles.get('Vulnerabilidad del agua subterránea', '—')}, la afectación podría extenderse hacia los "
        "receptores identificados en la sección 5.6.",
    ])

    rec = [r["receptor"] + (f": {r['descripcion']}" if r["descripcion"] else "")
           + (f", a {_n(r['distancia_m'])} m" if r.get("distancia_m") else "")
           + (f" hacia el {r['direccion']}" if r.get("direccion") else "")
           + (f", a {_n(r['profundidad_m'])} m de profundidad" if r.get("profundidad_m") else "") + "."
           for r in a["receptores"]]
    s["5.6"] = " ".join(rec) or PENDIENTE.format("receptores ecológicos del sitio")
    s["5.7"] = (" ".join(f"{x}." for x in a["pendientes"]) or "No se identificaron vacíos de información.") + \
        " La matriz de susceptibilidad es un tamizaje cualitativo y no sustituye una modelación numérica."

    concl = []
    if e["puntos_con_excedencia"]:
        concl.append(f"Se confirma afectación del suelo por {ind.lower()} en {e['puntos_con_excedencia']} de "
                     f"{e['puntos_evaluados']} puntos.")
    if not (d["horizontal"]["cerrada"] and d["vertical"]["cerrada"]):
        concl.append("La delimitación de la afectación no está cerrada; se requieren puntos adicionales "
                     "para completarla antes de dimensionar el saneamiento.")
    altas = [f for f, n in niveles.items() if n == "alta"]
    if altas:
        concl.append("Factores con susceptibilidad alta: " + ", ".join(x.lower() for x in altas) + ".")
    s["5.8"] = " ".join(concl) or PENDIENTE.format("conclusiones")
    return [{"numero": n, "titulo": SECCIONES[n], "texto": s[n]} for n in SECCIONES]


# ------------------------------------------------------------------ IA
SISTEMA = f"""Eres especialista en caracterización de sitios contaminados por hidrocarburos en México. Redactas el \
Capítulo 5, «Análisis de dispersión de hidrocarburos», del Informe de Caracterización que CAEPSA presentará a la \
autoridad ambiental.

Reglas estrictas:
- Usa exclusivamente los hechos del bloque <hechos>. No agregues cifras, distancias, fechas, profundidades, \
volúmenes ni datos del sitio que no estén ahí. Si un dato necesario falta, escribe «[DATO PENDIENTE: …]».
- Cita la fuente de cada dato del medio físico la primera vez que lo uses: «(fuente: …)».
- El único marco normativo es la {NORMA} (DOF {PUBLICACION_DOF}), y solo para LMP y excedencias. No cites otras \
normas, guías ni modelos externos.
- Los escenarios y rutas de migración son cualitativos y deben derivarse de la matriz de susceptibilidad y de la \
distribución observada; no hagas modelación numérica. Si hay tiempo de tránsito, preséntalo con su supuesto.
- Si la delimitación horizontal o vertical está abierta, dilo explícitamente en 5.3, 5.7 y 5.8.
- Redacción técnica en español, impersonal, en párrafos (sin viñetas). Puedes referir «Figura 5.1», «Tabla 5.1», \
«Tabla 5.2» y «Tabla 5.3», que se insertan automáticamente.
- Devuelve las ocho secciones 5.1 a 5.8 con los títulos indicados."""

ESQUEMA = {
    "type": "object",
    "properties": {"secciones": {"type": "array", "items": {
        "type": "object",
        "properties": {"numero": {"type": "string", "enum": list(SECCIONES)}, "texto": {"type": "string"}},
        "required": ["numero", "texto"], "additionalProperties": False}}},
    "required": ["secciones"],
    "additionalProperties": False,
}


class ErrorCapitulo(Exception):
    pass


def hechos(a: dict[str, Any]) -> dict[str, Any]:
    """Hechos para el redactor (sin la geometría de la envolvente)."""
    h = json.loads(json.dumps(a, default=str))
    h.get("espacial", {}).pop("envolvente", None)
    return h


def redactar_ia(a: dict[str, Any], plantilla: list[dict[str, str]], cliente: Any | None = None) -> list[dict[str, str]]:
    cliente = cliente or anthropic.Anthropic()
    mensaje = (f"<hechos>\n{json.dumps(hechos(a), ensure_ascii=False, indent=1)}\n</hechos>\n\n"
               "<titulos>\n" + "\n".join(f"{n} {t}" for n, t in SECCIONES.items()) + "\n</titulos>\n\n"
               "<borrador_base>\n" + "\n\n".join(f"{x['numero']} {x['titulo']}\n{x['texto']}" for x in plantilla)
               + "\n</borrador_base>\n\nRedacta el Capítulo 5 completo a partir de los hechos. Mejora la redacción "
               "técnica del borrador base sin agregar información que no esté en los hechos.")
    with cliente.beta.messages.stream(
        model=MODELO, max_tokens=32000, system=SISTEMA,
        messages=[{"role": "user", "content": mensaje}],
        output_config={"effort": "high", "format": {"type": "json_schema", "schema": ESQUEMA}},
        betas=["server-side-fallback-2026-07-01"], fallbacks="default",
    ) as stream:
        r = stream.get_final_message()
    if r.stop_reason == "refusal":
        raise ErrorCapitulo("El modelo declinó redactar el capítulo")
    if r.stop_reason == "max_tokens":
        raise ErrorCapitulo("La redacción excedió la longitud máxima de respuesta")
    try:
        datos = json.loads(next((b.text for b in r.content if b.type == "text"), ""))
    except json.JSONDecodeError as e:
        raise ErrorCapitulo(f"Respuesta del modelo no válida: {e}") from e
    por_num = {x["numero"]: x["texto"].strip() for x in datos["secciones"] if x["texto"].strip()}
    base = {x["numero"]: x["texto"] for x in plantilla}
    # Una sección que la IA omita conserva el texto de la plantilla.
    return [{"numero": n, "titulo": t, "texto": por_num.get(n, base[n])} for n, t in SECCIONES.items()]


# ------------------------------------------------------------------ verificación de cifras
_RE_NUM = re.compile(r"(?<![\w.])\d{1,3}(?:,\d{3})+(?:\.\d+)?|(?<![\w.])\d+(?:\.\d+)?")
_RE_EXCLUIR = re.compile(r"NOM[-\s]*\d{3}[-\w/]*|(?:Tabla|Figura|Cap[ií]tulo|apartado|secci[oó]n|Escenario)\s+"
                         r"\d+(?:\.\d+)*|\b5\.[1-8]\b|DOF\s+\d{2}/\d{2}/\d{4}|\[DATO PENDIENTE[^\]]*\]", re.I)


def _numeros(texto: str) -> list[tuple[str, float]]:
    limpio = _RE_EXCLUIR.sub(" ", texto)
    return [(m.group(), float(m.group().replace(",", ""))) for m in _RE_NUM.finditer(limpio)]


def cifras_no_respaldadas(secciones: list[dict[str, str]], a: dict[str, Any],
                          plantilla: list[dict[str, str]]) -> list[dict[str, str]]:
    """Cifras del texto que no aparecen en los hechos ni en la plantilla (tolerancia de redondeo)."""
    referencia = json.dumps(hechos(a), ensure_ascii=False) + "\n" + "\n".join(x["texto"] for x in plantilla)
    permitidos = {v for _, v in _numeros(referencia)}

    def respaldado(v: float) -> bool:
        if v.is_integer() and v <= 12:  # conteos pequeños, meses, números de escenario
            return True
        return any(abs(v - p) <= max(0.005 * abs(p), 0.051) or round(p) == v or round(p, 1) == v or round(p, 2) == v
                   for p in permitidos)

    out = []
    for s in secciones:
        for txt, v in _numeros(s["texto"]):
            if not respaldado(v):
                out.append({"seccion": s["numero"], "cifra": txt})
    return out


# ------------------------------------------------------------------ persistencia y orquestación
def generar(p: Proyecto, usar_ia: bool = True, cliente: Any | None = None) -> int:
    a = analizar(p)
    plantilla = redactar_plantilla(a)
    secciones, aviso, modelo, sospechosas = plantilla, None, None, []
    if usar_ia and a.get("disponible"):
        try:
            secciones = redactar_ia(a, plantilla, cliente)
            modelo = MODELO
            sospechosas = cifras_no_respaldadas(secciones, a, plantilla)
        except ErrorCapitulo as e:
            aviso = f"Redacción con IA no realizada; se usa la plantilla: {e}"
    with p.con:
        cur = p.con.execute(
            "INSERT INTO analisis_dispersion (usuario, metricas, capitulo, usa_ia, modelo, aviso, cifras_no_respaldadas)"
            " VALUES (?,?,?,?,?,?,?)",
            (p.usuario, json.dumps(a, ensure_ascii=False, default=str),
             json.dumps({"secciones": secciones, "plantilla": plantilla}, ensure_ascii=False),
             int(modelo is not None), modelo, aviso, json.dumps(sospechosas, ensure_ascii=False)))
        p.registrar_cambio("capitulo5", "analisis_dispersion", "generar", None,
                           {"id": cur.lastrowid, "usa_ia": modelo is not None, "cifras_no_respaldadas": len(sospechosas)})
    return int(cur.lastrowid or 0)


def listar(p: Proyecto) -> list[dict[str, Any]]:
    out = []
    for f in p.con.execute("SELECT * FROM analisis_dispersion ORDER BY id DESC"):
        r = dict(f)
        r["metricas"] = json.loads(r["metricas"])
        r["capitulo"] = json.loads(r["capitulo"])
        r["cifras_no_respaldadas"] = json.loads(r["cifras_no_respaldadas"] or "[]")
        out.append(r)
    return out


def obtener(p: Proyecto, analisis_id: int | None = None) -> dict[str, Any]:
    todos = listar(p)
    r = todos[0] if todos and analisis_id is None else next((x for x in todos if x["id"] == analisis_id), None)
    if r is None:
        raise KeyError("No hay capítulo generado" if analisis_id is None else f"No existe el análisis {analisis_id}")
    return r


# ------------------------------------------------------------------ Word
def a_word(p: Proyecto, analisis_id: int | None = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
    from docx.shared import Cm, Pt

    r = obtener(p, analisis_id)
    a = r["metricas"]
    marcar = {}
    for c in r["cifras_no_respaldadas"]:
        marcar.setdefault(c["seccion"], set()).add(c["cifra"])

    doc = Document()
    for lado in ("left_margin", "right_margin"):
        setattr(doc.sections[0], lado, Cm(2.5))
    doc.add_heading("5. ANÁLISIS DE DISPERSIÓN DE HIDROCARBUROS", level=1)

    def parrafo(texto: str, cifras: set[str]) -> None:
        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if not cifras:
            par.add_run(texto)
            return
        patron = re.compile("(" + "|".join(rf"(?<![\w.]){re.escape(c)}(?![\w])" for c in sorted(cifras, key=len,
                                                                                               reverse=True)) + ")")
        for i, trozo in enumerate(patron.split(texto)):
            run = par.add_run(trozo)
            if i % 2 == 1:
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    def tabla(titulo: str, encabezados: list[str], filas: list[list[Any]], nota: str = "") -> None:
        cap = doc.add_paragraph()
        cap.add_run(titulo).bold = True
        t = doc.add_table(rows=1, cols=len(encabezados))
        t.style = "Table Grid"
        for c, h in zip(t.rows[0].cells, encabezados):
            c.text = h
            c.paragraphs[0].runs[0].bold = True
        for f in filas:
            for c, v in zip(t.add_row().cells, f):
                c.text = "—" if v is None or v == "" else str(v)
        for fila in t.rows:
            for c in fila.cells:
                for par in c.paragraphs:
                    for run in par.runs:
                        run.font.size = Pt(8)
        if nota:
            doc.add_paragraph(nota).runs[0].font.size = Pt(8)

    for s in r["capitulo"]["secciones"]:
        doc.add_heading(f"{s['numero']} {s['titulo']}", level=2)
        for bloque in s["texto"].split("\n\n"):
            if bloque.strip():
                parrafo(bloque.strip(), marcar.get(s["numero"], set()))
        if not a.get("disponible"):
            continue
        if s["numero"] == "5.2":
            tabla("Tabla 5.1. Datos del medio físico y fuentes documentales.", ["Componente", "Dato", "Valor", "Fuente"],
                  [[mf.POR_CLAVE[k].grupo, d["etiqueta"], _con_unidad(k, d), d["fuente"]]
                   for k, d in a["medio_fisico"].items()])
        elif s["numero"] == "5.3":
            png = plano(a)
            if png:
                doc.add_picture(io.BytesIO(png), width=Cm(15))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.add_run("Figura 5.1. Distribución espacial de los puntos con excedencia de LMP.").bold = True
            tabla("Tabla 5.2. Afectación por punto de muestreo.",
                  ["Punto", "Zona", "Muestras", "Prof. máx. muestreada (m)", "Prof. máx. con excedencia (m)",
                   "Parámetros que exceden", "Delimitación vertical"],
                  [[x["punto"], x["zona"], x["muestras"], x["prof_max_muestreada_m"], x["prof_max_excedencia_m"],
                    ", ".join(x["parametros_excedidos"]), "Abierta" if x["vertical_abierta"] else "Cerrada"]
                   for x in a["puntos"]])
        elif s["numero"] == "5.4":
            tabla("Tabla 5.3. Matriz de susceptibilidad a la migración.", ["Factor", "Susceptibilidad", "Justificación"],
                  [[x["factor"], x["nivel"].capitalize(), x["justificacion"]] for x in a["matriz"]],
                  nota=a["criterio_tamizaje"])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
