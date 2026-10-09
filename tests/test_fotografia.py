"""Módulo 4 con imágenes generadas aquí y un cliente de API simulado (sin llamadas reales)."""
import io
import json
from types import SimpleNamespace as NS

import pytest
from docx import Document
from PIL import Image
from PIL.TiffImagePlugin import IFDRational

from caepsa.fotografia.clasificador import MAX_PALABRAS, contar_palabras, normalizar_pie
from caepsa.fotografia.metadatos import a_utm, leer_metadatos, punto_cercano
from caepsa.fotografia.registro import (
    actualizar_fotografia, clasificar_fotografias, diagnostico, listar_fotografias, registrar_fotografia)
from caepsa.fotografia.seccion import generar_word
from caepsa.proyectos import Proyecto


def _dms(valor):
    g = int(valor); m = int((valor - g) * 60); s = round(((valor - g) * 60 - m) * 60, 4)
    return (IFDRational(g, 1), IFDRational(m, 1), IFDRational(int(s * 10000), 10000))


def jpeg(color=(120, 90, 60), lat=None, lon=None, fecha=None):
    img = Image.new("RGB", (640, 480), color)
    exif = Image.Exif()
    if fecha:
        exif.get_ifd(0x8769)[36867] = fecha
    if lat is not None:
        gps = exif.get_ifd(0x8825)
        gps.update({1: "N" if lat >= 0 else "S", 2: _dms(abs(lat)), 3: "E" if lon >= 0 else "W", 4: _dms(abs(lon))})
    buf = io.BytesIO()
    img.save(buf, format="JPEG", exif=exif)
    return buf.getvalue()


def test_utm_contra_referencia_pyproj():
    # Valores de referencia calculados con pyproj (EPSG:4326 -> EPSG:32613/32614).
    x, y, z = a_utm(19.03, -102.1)
    assert z == 13 and x == pytest.approx(805273.279, abs=0.01) and y == pytest.approx(2106667.457, abs=0.01)
    x, y, z = a_utm(19.2, -101.95)
    assert z == 14 and x == pytest.approx(189777.497, abs=0.01)
    x, _, z = a_utm(19.2, -101.95, zona=13)  # zona forzada por el proyecto
    assert z == 13 and x == pytest.approx(820746.233, abs=0.01)


def test_exif_fecha_y_gps():
    md = leer_metadatos(jpeg(lat=19.03, lon=-102.1, fecha="2025:03:12 10:30:00"))
    assert md.fecha_toma == "2025-03-12 10:30:00"
    assert md.lat == pytest.approx(19.03, abs=1e-6) and md.lon == pytest.approx(-102.1, abs=1e-6)
    assert leer_metadatos(jpeg()).lat is None
    assert leer_metadatos(b"no es imagen").fecha_toma is None


def test_punto_cercano():
    puntos = [{"clave": "A", "utm_x": 0, "utm_y": 0}, {"clave": "B", "utm_x": 20, "utm_y": 0},
              {"clave": "C", "utm_x": None, "utm_y": None}]
    assert punto_cercano(15, 0, puntos) == ("B", 5)
    assert punto_cercano(500, 0, puntos) is None


def test_pie_normalizado():
    assert normalizar_pie("  vista general   del área afectada. ") == "Vista general del área afectada"
    assert contar_palabras("Toma de muestra de suelo superficial") == 6


class VisionFalsa:
    """Devuelve respuestas JSON en orden; registra cada solicitud."""

    def __init__(self, respuestas):
        self.respuestas = list(respuestas)
        self.solicitudes = []
        self.beta = NS(messages=NS(create=self._create))

    def _create(self, **kw):
        self.solicitudes.append(kw)
        r = self.respuestas.pop(0)
        if r == "refusal":
            return NS(stop_reason="refusal", content=[])
        return NS(stop_reason="end_turn", content=[NS(type="text", text=json.dumps(r))])


def _resp(cat, pie, conf="alta", obs=""):
    return {"categoria": cat, "pie": pie, "confianza": conf, "observacion": obs}


def test_flujo_registro_clasificacion_y_seccion():
    p = Proyecto.crear("Fotos")
    p.actualizar_ficha({"utm_zona": "13Q", "sitio": "Km 176+500 Autopista Siglo XXI"})
    x, y, _ = a_utm(19.03, -102.1, 13)
    p.guardar_punto("Z1-P03", zona="Zona 1", utm_x=x + 8, utm_y=y - 5)

    con_gps = registrar_fotografia(p, "IMG_001.jpg", jpeg(lat=19.03, lon=-102.1, fecha="2025:03:12 10:30:00"))
    sin_gps = registrar_fotografia(p, "IMG_002.jpg", jpeg(color=(10, 120, 10)))
    rechazo = registrar_fotografia(p, "IMG_003.jpg", jpeg(color=(200, 200, 200)))
    assert registrar_fotografia(p, "copia.jpg", jpeg(lat=19.03, lon=-102.1, fecha="2025:03:12 10:30:00")) == con_gps

    foto = next(f for f in listar_fotografias(p) if f["id"] == con_gps)
    assert foto["punto"] == "Z1-P03" and foto["punto_origen"] == "gps"

    largo = " ".join(["palabra"] * 20)
    # Orden de proceso: sin fecha primero (IMG_002, IMG_003) y luego IMG_001.
    cliente = VisionFalsa([
        _resp("reconocimiento", largo), _resp("reconocimiento", "Vista general del área afectada por el derrame"),
        "refusal",
        _resp("muestreo", "Toma de muestra de suelo superficial con pala de acero inoxidable."),
    ])
    res = clasificar_fotografias(p, cliente=cliente)
    assert res["procesadas"] == 2 and len(res["errores"]) == 1 and "IMG_003" in res["errores"][0]
    # La primera foto se reintentó por exceso de palabras; el punto asociado por GPS llega como contexto.
    assert "Redúcela" in cliente.solicitudes[1]["messages"][0]["content"][1]["text"]
    assert "Z1-P03 (Zona 1)" in cliente.solicitudes[3]["messages"][0]["content"][1]["text"]
    assert cliente.solicitudes[0]["output_config"]["format"]["type"] == "json_schema"

    fotos = {f["id"]: f for f in listar_fotografias(p)}
    assert fotos[con_gps]["pie"] == "Toma de muestra de suelo superficial con pala de acero inoxidable"
    assert fotos[sin_gps]["categoria"] == "reconocimiento" and fotos[sin_gps]["pie_origen"] == "ia"
    assert contar_palabras(fotos[sin_gps]["pie"]) <= MAX_PALABRAS
    assert fotos[rechazo]["categoria"] is None

    # Reclasificación manual: queda marcada y la IA no la pisa en una nueva pasada.
    actualizar_fotografia(p, con_gps, categoria="excavaciones", revisada=1)
    clasificar_fotografias(p, cliente=VisionFalsa([_resp("muestreo", "Otra cosa")]), ids=[con_gps])
    f = next(x for x in listar_fotografias(p) if x["id"] == con_gps)
    assert f["categoria"] == "excavaciones" and f["categoria_origen"] == "manual"
    with pytest.raises(ValueError):
        actualizar_fotografia(p, con_gps, categoria="inventada")

    d = diagnostico(p)
    assert d["sin_clasificar"] == [rechazo] and d["por_categoria"]["reconocimiento"] == 1

    contenido, avisos = generar_word(p)
    doc = Document(io.BytesIO(contenido))
    textos = [par.text for t in doc.tables for c in t._cells for par in c.paragraphs if par.text]
    # Orden del informe: reconocimiento antes que excavaciones; numeración continua.
    assert textos[0].startswith("Fotografía 1. Vista general") and textos[1].startswith("Fotografía 2. Otra cosa")
    assert len(doc.inline_shapes) == 2
    assert any("IMG_003.jpg" in a for a in avisos)


def test_herramientas_fotograficas_del_asistente_y_tablero():
    from caepsa.asistente.herramientas import Contexto, ejecutar
    from caepsa.tablero import estado

    ctx = Contexto(usuario="qa", cliente=VisionFalsa([_resp("cierre", "Retiro de materiales al concluir los trabajos")]))
    ejecutar(ctx, "crear_proyecto", {"nombre": "Fotos asistente"})
    p = ctx.abrir()
    _, alertas = estado(p)
    assert any("No se han cargado fotografías" in a.mensaje for a in alertas)

    foto = registrar_fotografia(p, "cierre.jpg", jpeg())
    salida, error = ejecutar(ctx, "clasificar_fotografias", {})
    assert not error and json.loads(salida)["procesadas"] == 1
    lista = json.loads(ejecutar(ctx, "listar_fotografias", {})[0])
    assert lista[0]["categoria"] == "cierre" and lista[0]["palabras"] == 7 and lista[0]["revisada"] == 0

    _, error = ejecutar(ctx, "actualizar_fotografia", {"id": foto, "pie": " ".join(["x"] * 16)})
    assert error  # más de 15 palabras
    salida, error = ejecutar(ctx, "actualizar_fotografia", {"id": foto, "revisada": True})
    assert not error and json.loads(salida)["cambios"]["revisada"] == 1

    avance, alertas = estado(p)
    assert avance[4] == 1.0
    assert any("Sin evidencia fotográfica de: Reconocimiento del sitio, Muestreo de suelo" in a.mensaje for a in alertas)
    salida, error = ejecutar(ctx, "generar_seccion_fotografica", {"titulo": "Anexo fotográfico"})
    assert not error and ctx.archivos_generados[-1].name.endswith("_seccion_fotografica.docx")
