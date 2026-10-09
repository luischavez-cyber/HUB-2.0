"""Módulo 5 con borradores SINTÉTICOS con errores sembrados y un cliente de API simulado."""
import io
import json
from datetime import date
from types import SimpleNamespace as NS

import pytest
from docx import Document as NuevoDocx
from openpyxl import load_workbook
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle

from caepsa.laboratorio.captura import importar_csv
from caepsa.proyectos import Proyecto
from caepsa.revision.dictamen import a_excel, a_word
from caepsa.revision.documento import leer
from caepsa.revision.indice import Hallazgo, calcular_ict
from caepsa.revision.registro import actualizar_hallazgo, hallazgos_de, listar_revisiones, revisar_borrador
from caepsa.revision.revisor_ia import evidencia_localizada
from caepsa.revision.verificaciones import verificar

HOY = date(2026, 10, 9)


def proyecto():
    p = Proyecto.crear("Revisión")
    p.actualizar_ficha({"usos_suelo": ["agricola"], "producto": "Gasolinas", "utm_zona": "13Q"})
    errores = importar_csv(p, (
        "muestra,punto,zona,profundidad_m,utm_x,utm_y,fecha_muestreo,parametro,valor\n"
        "Z1-P01-M1,Z1-P01,Zona 1,0.30,801234.50,2108765.20,12/03/2025,HFL,512.3\n"
        "Z1-P01-M1,Z1-P01,Zona 1,0.30,801234.50,2108765.20,12/03/2025,BENCENO,<0.002\n"
        "Z2-P05-M1,Z2-P05,Zona 2,0.30,801300.00,2108800.00,12/03/2025,HFL,85.1\n"
        "Z2-P05-M1,Z2-P05,Zona 2,0.30,801300.00,2108800.00,12/03/2025,BENCENO,<0.002\n"))
    assert errores == []
    return p


def borrador_con_errores() -> bytes:
    d = NuevoDocx()
    d.add_heading("1. Antecedentes", level=1)
    d.add_paragraph("El derrame ocurrió en el Km 176+500 de la Autopista Siglo XXI, Nueva Italia, Michoacán.")
    d.add_heading("2. Marco legal", level=1)
    d.add_paragraph("La caracterización se realizó conforme a la NOM-138-SEMARNAT/SS-2003 y la NOM-052-SEMARNAT-2005.")
    d.add_heading("3. Metodología de muestreo", level=1)
    d.add_paragraph("Se tomaron 42 muestras en 3 puntos de muestreo el 31/02/2025. El uso de suelo es industrial.")
    d.add_heading("4. Resultados", level=1)
    t = d.add_table(rows=4, cols=6)
    for fila, valores in zip(t.rows, [
        ["Muestra", "X (UTM)", "Y (UTM)", "Fracción ligera", "Benceno", "Dictamen"],
        ["LMP", "", "", "500", "6", ""],
        ["Z1-P01-M1", "801234.50", "2108765.20", "450.0", "<0.002", "Cumple"],
        ["Z2-P05-M1", "2108800", "801300", "85.1", "<0.002", "Cumple"],
    ]):
        for c, v in zip(fila.cells, valores):
            c.text = v
    d.add_paragraph("Todas las muestras cumplen con los límites máximos permisibles, expresados en ppm.")
    d.add_heading("5. Conclusiones", level=1)
    d.add_paragraph("El sitio no requiere acciones adicionales.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_lectura_docx_conserva_secciones_y_tablas():
    doc = leer(borrador_con_errores(), "b.docx")
    assert [b.texto for b in doc.titulos][:2] == ["1. Antecedentes", "2. Marco legal"]
    assert doc.tablas[0].filas[2][0] == "Z1-P01-M1" and doc.tablas[0].seccion == "4. Resultados"
    assert "[Tabla 1 (sección «4. Resultados»)]" in doc.como_texto_marcado()


def test_lectura_pdf():
    buf = io.BytesIO()
    st = getSampleStyleSheet()
    SimpleDocTemplate(buf, pagesize=letter).build([
        Paragraph("4. RESULTADOS", st["Heading1"]),
        Paragraph("Se analizaron 2 muestras conforme a la NOM-138-SEMARNAT/SSA1-2012.", st["Normal"]),
        Table([["Muestra", "Fracción ligera"], ["Z1-P01-M1", "512.3"]],
              style=TableStyle([("GRID", (0, 0), (-1, -1), 0.5, "black")])),
        Table([["Sin bordes", "Z9-P09-M9"]]),
    ])
    doc = leer(buf.getvalue(), "b.pdf")
    assert doc.paginas == 1 and doc.titulos[0].texto == "4. RESULTADOS"
    assert len(doc.tablas) == 1 and doc.tablas[0].filas[1] == ["Z1-P01-M1", "512.3"]
    parrafos = " ".join(b.texto for b in doc.bloques if b.tipo == "parrafo")
    assert "Z1-P01-M1" not in parrafos  # la tabla con bordes no se duplica como texto
    assert "Z9-P09-M9" in parrafos  # la tabla sin bordes no se pierde: queda como texto


def test_verificaciones_detectan_errores_sembrados():
    p = proyecto()
    hs = verificar(leer(borrador_con_errores(), "b.docx"), p, HOY)
    txt = "\n".join(f"{h.criticidad}|{h.categoria}|{h.descripcion}" for h in hs)
    esperados = [
        "critica|vacio_regulatorio|El informe no cita la NOM-138-SEMARNAT/SSA1-2012",
        "alta|referencia_normativa|Se cita la NOM-138-SEMARNAT/SS-2003",
        "info|referencia_normativa|Se citan normas distintas a la NOM-138: NOM-052-SEMARNAT-2005",
        "critica|vacio_regulatorio|LMP de Fracción ligera citado: 500 mg/kg",
        "alta|inconsistencia_laboratorio|Z1-P01-M1, Fracción ligera: el informe reporta 450.0 y el laboratorio 512.3",
        "critica|contradictorio|La tabla dictamina «Cumple» para Z1-P01-M1",
        "alta|fuera_de_rango|Coordenadas UTM fuera de rango",
        "critica|contradictorio|El texto afirma cumplimiento, pero 1 muestra(s) exceden",
        "alta|inconsistencia_tabla_texto|El informe refiere uso de suelo «industrial»",
        "alta|inconsistencia_laboratorio|El informe indica 42 muestras; la base del proyecto registra 2",
        "alta|inconsistencia_laboratorio|El informe indica 3 puntos de muestreo; la base del proyecto registra 2",
        "media|fechas|Fecha inválida: «31/02/2025»",
        "media|fechas|Las fechas de muestreo registradas por el laboratorio no aparecen en el informe: 12/03/2025",
        "baja|ambiguo|Se usan «ppm»",
        "alta|estructura|No se identifica la sección «Plan de saneamiento»",  # hay excedencias
        "media|estructura|No se identifica la sección «Identificación de la pluma contaminante»",
        "media|insuficiente|No se menciona la cadena de custodia",
        "media|insuficiente|No se documenta la acreditación del laboratorio",
    ]
    for e in esperados:
        assert e in txt, f"No se detectó: {e}\n---\n{txt}"
    # Sin falsos positivos en lo que sí es correcto.
    assert "Z2-P05-M1, Fracción ligera" not in txt
    assert "LMP de Benceno" not in txt
    assert "sección «Antecedentes»" not in txt and "sección «Resultados»" not in txt


def test_ict_determinista_y_descartados():
    hs = [Hallazgo("automatica", "contradictorio", "critica", "x", "a"),
          Hallazgo("ia", "ambiguo", "alta", "x", "b", evidencia="no existe", evidencia_localizada=False)]
    r = calcular_ict(hs)
    assert r["subindices"] == {"consistencia": 80.0, "normativo": 100.0, "completitud": 100.0, "claridad": 95.0}
    assert r["ict"] == pytest.approx(80 * .35 + 100 * .30 + 100 * .20 + 95 * .15, abs=0.05)
    hs[0].estado = "descartado"
    assert calcular_ict(hs)["subindices"]["consistencia"] == 100.0
    assert calcular_ict([])["interpretacion"].startswith("Apto")
    with pytest.raises(ValueError):
        Hallazgo("ia", "inventada", "alta", "x", "y")


def test_evidencia_localizada():
    texto = "el derrame ocurrio en el km 176+500 de la autopista siglo xxi"
    assert evidencia_localizada("El derrame ocurrió en el Km 176+500", texto) is True
    assert evidencia_localizada("El derrame … Autopista Siglo XXI", texto) is True
    assert evidencia_localizada("Se instalaron 4 pozos de monitoreo", texto) is False
    assert evidencia_localizada("", texto) is None


class _CM:
    def __init__(self, r): self.r = r
    def __enter__(self): return NS(get_final_message=lambda: self.r)
    def __exit__(self, *a): return False


class RevisorFalsoCM:
    """Simula client.beta.messages.stream(...) como administrador de contexto."""

    def __init__(self, datos, stop="end_turn"):
        self.datos, self.stop, self.solicitudes = datos, stop, []
        self.beta = NS(messages=NS(stream=self._stream))

    def _stream(self, **kw):
        self.solicitudes.append(kw)
        return _CM(NS(stop_reason=self.stop, content=[NS(type="text", text=json.dumps(self.datos))]))


def test_revision_completa_con_ia_historial_y_exportacion():
    p = proyecto()
    p.guardar_documento("borrador", "borrador_v1.docx", borrador_con_errores())
    doc_id = p.documentos("borrador")[0]["id"]
    ia = RevisorFalsoCM({
        "resumen": "El informe presenta contradicciones graves entre resultados y conclusiones.",
        "fortalezas": ["Ubicación del evento claramente identificada"],
        "hallazgos": [
            {"categoria": "insuficiente", "criticidad": "alta", "ubicacion": "5. Conclusiones",
             "evidencia": "El sitio no requiere acciones adicionales", "descripcion": "Conclusión sin sustento.",
             "riesgo_regulatorio": "Prevención", "recomendacion": "Sustentar con resultados.",
             "referencia_nom138": "NOM-138 (verificar apartado)"},
            {"categoria": "ambiguo", "criticidad": "media", "ubicacion": "3",
             "evidencia": "Se instalaron cuatro pozos de monitoreo", "descripcion": "Cita inexistente.",
             "riesgo_regulatorio": "", "recomendacion": "Verificar.", "referencia_nom138": ""},
        ]})
    rev_id = revisar_borrador(p, doc_id, usar_ia=True, cliente=ia, hoy=HOY)
    kw = ia.solicitudes[0]
    assert kw["model"] == "claude-opus-5-5" and kw["output_config"]["format"]["type"] == "json_schema"
    mensaje = kw["messages"][0]["content"]
    assert "Excede: Z1-P01-M1" in mensaje and "LMP de Fracción ligera citado" in mensaje  # contexto y no-repetición

    hs = hallazgos_de(p, rev_id)
    ia_hs = [h for h in hs if h.origen == "ia"]
    assert [h.evidencia_localizada for h in ia_hs] == [True, False]
    rev = listar_revisiones(p)[0]
    assert rev["usa_ia"] == 1 and rev["ict"] < 50 and rev["fortalezas"] == ["Ubicación del evento claramente identificada"]

    critico = next(h for h in hs if h.criticidad == "critica")
    with pytest.raises(ValueError):
        actualizar_hallazgo(p, critico.id, "descartado")  # requiere justificación
    actualizar_hallazgo(p, critico.id, "descartado", "Falso positivo validado por el responsable técnico")
    assert listar_revisiones(p)[0]["ict_vigente"] > rev["ict"]
    assert p.bitacora()[0]["entidad"] == "hallazgo"

    doc = NuevoDocx(io.BytesIO(a_word(p, rev_id)))
    assert any(par.text.startswith("Índice de Calidad Técnica:") for par in doc.paragraphs)
    assert "cita no localizada" in "\n".join(c.text for t in doc.tables for c in t._cells)
    wb = load_workbook(io.BytesIO(a_excel(p, rev_id)))
    assert wb["Hallazgos"].max_row == len(hs) + 1 and wb["Índice"]["A1"].value == "Documento"


def test_sin_ia_o_con_rechazo_queda_aviso():
    p = proyecto()
    p.guardar_documento("borrador", "b.docx", borrador_con_errores())
    doc_id = p.documentos("borrador")[0]["id"]
    rev_id = revisar_borrador(p, doc_id, usar_ia=True, cliente=RevisorFalsoCM({}, stop="refusal"), hoy=HOY)
    rev = listar_revisiones(p)[0]
    assert rev["id"] == rev_id and rev["usa_ia"] == 0 and "declinó" in rev["aviso"]
    assert all(h.origen == "automatica" for h in hallazgos_de(p, rev_id))


def test_herramientas_del_asistente_y_tablero():
    from caepsa.asistente.herramientas import Contexto, ejecutar
    from caepsa.tablero import estado

    p = proyecto()
    ctx = Contexto(usuario="qa", proyecto=p.clave)
    p.guardar_documento("borrador", "borrador_v1.docx", borrador_con_errores())
    _, alertas = estado(p)
    assert any("borradores cargados sin revisión" in a.mensaje for a in alertas)

    doc_id = json.loads(ejecutar(ctx, "listar_borradores", {})[0])[0]["documento_id"]
    salida, error = ejecutar(ctx, "revisar_informe", {"documento_id": doc_id, "usar_ia": False})
    r = json.loads(salida)
    assert not error and r["hallazgos_por_criticidad"]["critica"] >= 4 and r["ict"] < 50

    criticos = json.loads(ejecutar(ctx, "consultar_revision", {"criticidad": "critica"})[0])["hallazgos"]
    assert criticos and all(h["criticidad"] == "critica" for h in criticos)
    _, error = ejecutar(ctx, "actualizar_hallazgo", {"id": criticos[0]["id"], "estado": "descartado"})
    assert error  # sin justificación
    assert not ejecutar(ctx, "actualizar_hallazgo", {"id": criticos[0]["id"], "estado": "atendido"})[1]

    avance, alertas = estado(p)
    assert 0 < avance[5] < 0.5
    assert any(a.modulo == 5 and "crítico(s) abiertos" in a.mensaje for a in alertas)

    for formato in ("word", "excel"):
        salida, error = ejecutar(ctx, "exportar_revision", {"formato": formato})
        assert not error and ctx.archivos_generados[-1].exists()
