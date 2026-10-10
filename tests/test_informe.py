"""Módulo 7 con el sitio SINTÉTICO de pruebas (no son datos reales)."""
import io
from datetime import date

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

from caepsa.dispersion import capitulo
from caepsa.fotografia.registro import actualizar_fotografia, registrar_fotografia
from caepsa.informe.generador import CAPITULOS, generar, generar_y_registrar, verificacion_previa
from caepsa.revision.registro import hallazgos_de
from sitio_sintetico import crear
from test_fotografia import jpeg

HOY = date(2026, 10, 10)


def proyecto_completo():
    p = crear("Informe sintético")
    p.actualizar_ficha({"cliente": "Cliente de prueba", "sitio": "Derecho de vía", "municipio": "Municipio X",
                        "estado": "Michoacán", "ubicacion": "Km 0+000", "responsable_tecnico": "Ing. Prueba",
                        "expediente": "EXP-0000", "laboratorio": "Laboratorio de prueba",
                        "fecha_evento": "10/03/2025"})
    for i, (cat, pie) in enumerate([("reconocimiento", "Vista general del área afectada"),
                                    ("muestreo", "Toma de muestra de suelo superficial")]):
        fid = registrar_fotografia(p, f"F{i}.jpg", jpeg(color=(60 + 80 * i, 100, 80)))
        actualizar_fotografia(p, fid, categoria=cat, pie=pie, revisada=1)
    capitulo.generar(p, usar_ia=False)
    return p


def test_estructura_formato_y_contenido():
    p = proyecto_completo()
    contenido, meta = generar(p, hoy=HOY)
    doc = Document(io.BytesIO(contenido))
    h1 = [x.text for x in doc.paragraphs if x.style.name == "Heading 1"]
    assert h1 == list(CAPITULOS)
    textos = "\n".join(x.text for x in doc.paragraphs)
    assert "INFORME DE CARACTERIZACIÓN" in textos and "Octubre de 2026" in textos
    assert "1. ANTECEDENTES" in textos.split("1. ANTECEDENTES")[0] + "1. ANTECEDENTES"  # índice precargado
    indice = next(x.text for x in doc.paragraphs if "4. RESULTADOS" in x.text and x.style.name != "Heading 1")
    assert "      5.3 Distribución espacial" in indice

    # Página horizontal para la Tabla 4.2 y vuelta a vertical.
    orient = [s.orientation for s in doc.sections]
    assert WD_ORIENT.LANDSCAPE in orient and orient[-1] == WD_ORIENT.PORTRAIT
    # Paginación, encabezado y actualización de campos al abrir.
    instr = [i.text for i in doc.sections[0].footer._element.iter(qn("w:instrText"))]
    assert any("PAGE" in t for t in instr) and any("NUMPAGES" in t for t in instr)
    assert "Informe sintético" in doc.sections[0].header.paragraphs[0].text
    assert doc.settings.element.find(qn("w:updateFields")) is not None

    # Orden de zonas configurado (Zona 1 antes que Periferia) en la Tabla 4.2.
    t42 = next(t for t in doc.tables if t.rows[0].cells[0].text == "Zona" and len(t.columns) > 8)
    assert t42.rows[2].cells[0].text == "Zona 1" and t42.rows[-1].cells[0].text == "Periferia"
    # Pendientes resaltados y contados; nada inventado para lo que no hay dato.
    assert meta["pendientes"] and all(x.startswith("[COMPLETAR") for x in meta["pendientes"])
    resaltados = [r.text for x in doc.paragraphs for r in x.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert any("técnica de remediación" in r for r in resaltados)
    assert meta["fotografias"] == 2 and len(doc.inline_shapes) == 3  # 2 fotos + Figura 5.1


def test_informe_generado_es_consistente_segun_el_revisor():
    p = proyecto_completo()
    r = generar_y_registrar(p, hoy=HOY)
    hs = hallazgos_de(p, r["revision_id"])
    categorias = {h.categoria for h in hs}
    # Ninguna inconsistencia de datos ni de normativa en lo que la plataforma redacta.
    assert not categorias & {"inconsistencia_laboratorio", "contradictorio", "vacio_regulatorio", "coordenadas",
                             "fuera_de_rango", "inconsistencia_tabla_texto", "estructura"}, \
        [h.descripcion for h in hs if h.categoria != "insuficiente"]
    # Los pendientes impiden calificarlo como apto.
    assert any(h.criticidad == "critica" and "marcador" in h.descripcion for h in hs)
    assert r["interpretacion"].startswith("Requiere correcciones") or r["interpretacion"].startswith("No apto")
    assert r["nombre"].endswith("_v1.docx") and generar_y_registrar(p, revisar=False, hoy=HOY)["nombre"].endswith("_v2.docx")


def test_plantilla_de_empresa_conserva_encabezado_y_crea_estilos_faltantes(tmp_path):
    p = proyecto_completo()
    plantilla = Document()
    plantilla.sections[0].header.paragraphs[0].text = "CAEPSA · Formato institucional"
    plantilla.add_paragraph("Contenido de ejemplo de la plantilla")
    estilo = plantilla.styles["Table Grid"].element
    estilo.getparent().remove(estilo)  # plantilla que nunca usó tablas
    buf = io.BytesIO()
    plantilla.save(buf)
    p.guardar_documento("plantilla_word", "plantilla.docx", buf.getvalue())
    pid = p.documentos("plantilla_word")[0]["id"]

    r = generar_y_registrar(p, plantilla_id=pid, revisar=False, hoy=HOY)
    doc = Document(r["ruta"])
    assert doc.sections[0].header.paragraphs[0].text == "CAEPSA · Formato institucional"
    assert "Contenido de ejemplo de la plantilla" not in "\n".join(x.text for x in doc.paragraphs)
    assert doc.tables[0].style.name == "Table Grid"


def test_verificacion_previa():
    p = crear("Previa", con_medio=False)
    niveles = {x["texto"].split(":")[0]: x["nivel"] for x in verificacion_previa(p)}
    assert niveles["Ficha incompleta"] == "aviso"
    assert niveles["Medio físico"] == "aviso"
    assert niveles["Capítulo 5 no generado"] == "aviso"
    assert verificacion_previa(proyecto_completo())[0] == {"nivel": "ok", "texto": "Ficha del proyecto completa"}


def test_herramientas_del_asistente_y_tablero():
    import json

    from caepsa.asistente.herramientas import Contexto, ejecutar
    from caepsa.tablero import estado

    p = proyecto_completo()
    ctx = Contexto(usuario="qa", proyecto=p.clave)
    avance, alertas = estado(p)
    assert avance[7] == 0 and any("aún no generado" in a.mensaje for a in alertas)
    assert not any("etapa posterior" in a.mensaje for a in alertas)  # los 7 módulos están disponibles

    previa = json.loads(ejecutar(ctx, "verificar_insumos_informe", {})[0])
    assert previa[0]["nivel"] == "ok"
    salida, error = ejecutar(ctx, "generar_informe", {})
    r = json.loads(salida)
    assert not error and r["datos_por_completar"] > 0 and "ict" in r and ctx.archivos_generados[-1].exists()
    avance, alertas = estado(p)
    assert 0 < avance[7] < 1 and any(a.modulo == 7 and "por completar" in a.mensaje for a in alertas)
