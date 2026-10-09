from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from caepsa.exportacion.tablas import a_excel, a_word
from caepsa.normativa.evaluacion import evaluar
from caepsa.normativa.nom138 import UsoSuelo

FILAS = [
    {"muestra": "M1", "punto": "P1", "zona": "Zona 1", "parametro": "HFL", "valor": 512.3, "menor_que": 0,
     "profundidad_m": "0.30", "utm_x": 801234.5, "utm_y": 2108765.2},
    {"muestra": "M2", "punto": "P2", "zona": "Periferia", "parametro": "HFL", "valor": 10, "menor_que": 1,
     "profundidad_m": "0.30", "utm_x": 801500.0, "utm_y": 2109001.0},
]


def test_excel_resalta_excedencias():
    ms = evaluar(FILAS, [UsoSuelo.AGRICOLA], None)
    wb = load_workbook(BytesIO(a_excel(ms, [UsoSuelo.AGRICOLA], {"nombre": "X"})))
    ws = wb["Cumplimiento NOM-138"]
    assert ws.cell(row=6, column=7).value == 200  # fila LMP
    assert ws.cell(row=7, column=7).value == "512.3"
    assert ws.cell(row=7, column=7).fill.fgColor.rgb.endswith("F8D7DA")
    assert ws.cell(row=8, column=7).value == "<10"
    assert {"Detalle", "Resumen"} <= set(wb.sheetnames)


def test_word_tabla():
    ms = evaluar(FILAS, [UsoSuelo.AGRICOLA, UsoSuelo.RESIDENCIAL], None)
    doc = Document(BytesIO(a_word(ms, [UsoSuelo.AGRICOLA, UsoSuelo.RESIDENCIAL], {"laboratorio": "Lab"})))
    t = doc.tables[0]
    assert t.rows[2].cells[6].text == "512.3"
    assert any("Uso de suelo mixto" in p.text for p in doc.paragraphs)
