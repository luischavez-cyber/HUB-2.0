"""Módulo 6 con un sitio SINTÉTICO (tests/sitio_sintetico.py) y un cliente de API simulado."""
import io
import json
from types import SimpleNamespace as NS

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from caepsa.dispersion import capitulo, medio_fisico as mf
from caepsa.dispersion.analisis import analizar, area_poligono, envolvente_convexa, profundidad, rumbo
from sitio_sintetico import crear


def test_geometria_basica():
    cuadrado = envolvente_convexa([(0, 0), (10, 0), (10, 10), (0, 10), (5, 5)])
    assert len(cuadrado) == 4 and area_poligono(cuadrado) == 100
    assert area_poligono(envolvente_convexa([(0, 0), (1, 1), (2, 2)])) == 0  # alineados
    assert rumbo(0) == "N" and rumbo(156) == "SE" and rumbo(350) == "N" and rumbo(225) == "SO"
    assert profundidad("0.0-0.3 m") == 0.3 and profundidad("1,5") == 1.5 and profundidad(None) is None


def test_medio_fisico_validacion_fuente_y_bitacora():
    p = crear("MF", con_medio=False)
    with pytest.raises(ValueError):
        mf.guardar(p, {"pendiente_pct": {"valor": 150, "fuente": "x"}})
    with pytest.raises(ValueError):
        mf.guardar(p, {"textura": {"valor": "barro", "fuente": "x"}})
    with pytest.raises(ValueError):
        mf.guardar(p, {"campo_inventado": {"valor": 1}})
    assert mf.guardar(p, {"pendiente_pct": {"valor": "3.5", "fuente": ""}}) == ["pendiente_pct"]
    datos = mf.leer(p)
    assert datos["pendiente_pct"]["valor"] == 3.5 and mf.valor(datos, "pendiente_pct") is None  # sin fuente
    falt = mf.faltantes(datos)
    assert any(c.clave == "pendiente_pct" for c in falt["sin_fuente"])
    assert any(c.clave == "textura" for c in falt["sin_valor"])
    assert p.bitacora()[0]["entidad"] == "medio_fisico"


def test_analisis_sitio_sintetico():
    a = analizar(crear())
    e, d = a["espacial"], a["delimitacion"]
    assert a["indicador"]["clave"] == "HFL" and a["origen"]["tipo"] == "documentado"
    assert e["puntos_con_excedencia"] == 5 and e["puntos_evaluados"] == 8
    assert e["extension_max_m"] == pytest.approx(76.2, abs=0.1)
    assert e["direccion_dispersion"]["rumbo"] == "SE" and e["coincidencia_pendiente"]["valoracion"] == "coincide"
    assert e["area_afectada_m2"] == pytest.approx(865.0, abs=0.5)
    assert e["prof_max_afectacion_m"] == 2.0
    assert e["volumen_estimado_m3"]["conservador"] == pytest.approx(1730.0, abs=1)
    assert not d["horizontal"]["cerrada"] and "PER-P06" in d["horizontal"]["motivo"]
    assert not d["vertical"]["cerrada"] and "Z1-P01 (2 m)" in d["vertical"]["motivo"]
    niveles = {x["factor"]: x["nivel"] for x in a["matriz"]}
    assert niveles == {"Infiltración y lixiviación vertical": "alta", "Vulnerabilidad del agua subterránea": "alta",
                       "Arrastre por escurrimiento superficial": "media",
                       "Afectación de cuerpos de agua superficiales": "alta", "Movilidad del contaminante": "alta"}
    assert a["transito_vertical"]["dias"] == pytest.approx(0.9, abs=0.05)
    assert "elevacion_msnm" not in a["medio_fisico"]  # sin fuente: no se usa
    assert any("sin fuente documentada (no se usa): Elevación" in x for x in a["pendientes"])


def test_analisis_sin_medio_fisico_no_inventa():
    a = analizar(crear("Sin medio", con_medio=False))
    assert a["origen"]["tipo"] == "punto de mayor concentración" and a["origen"]["punto"] == "Z1-P01"
    assert {x["nivel"] for x in a["matriz"]} - {"alta"} == {"indeterminado"}  # solo movilidad es determinable
    assert a["transito_vertical"] is None and a["receptores"] == []
    plantilla = capitulo.redactar_plantilla(a)
    assert len(plantilla) == 8 and "[DATO PENDIENTE: pendiente del terreno]" in plantilla[1]["texto"]


def test_plantilla_con_datos():
    secciones = {s["numero"]: s["texto"] for s in capitulo.redactar_plantilla(analizar(crear()))}
    assert "X 801200.00, Y 2108700.00" in secciones["5.3"]
    assert "Hidrología subterránea: profundidad al nivel freático: 8 m. Fuente: Censo de pozos CONAGUA." in secciones["5.2"]
    assert secciones["5.1"].startswith("En este capítulo se analiza el comportamiento en el suelo del producto derramado (gasolinas), con base")
    assert "Delimitación horizontal: abierta" in secciones["5.3"]
    assert "no está cerrada" in secciones["5.8"]


class _CM:
    def __init__(self, r): self.r = r
    def __enter__(self): return NS(get_final_message=lambda: self.r)
    def __exit__(self, *a): return False


class RedactorFalso:
    def __init__(self, secciones=None, stop="end_turn"):
        self.secciones, self.stop, self.solicitudes = secciones, stop, []
        self.beta = NS(messages=NS(stream=self._stream))

    def _stream(self, **kw):
        self.solicitudes.append(kw)
        texto = json.dumps({"secciones": self.secciones or []})
        return _CM(NS(stop_reason=self.stop, content=[NS(type="text", text=texto)]))


def test_capitulo_con_ia_verifica_cifras_y_exporta_word():
    p = crear()
    redactor = RedactorFalso([
        {"numero": "5.1", "texto": "Se analizan 15 muestras de 8 puntos conforme a la NOM-138-SEMARNAT/SSA1-2012."},
        {"numero": "5.3", "texto": "La excedencia más alejada está a 76.2 m del origen, en un área de 865 m². "
                                   "El canal se ubica a 350 m (dato inventado)."},
    ])
    aid = capitulo.generar(p, usar_ia=True, cliente=redactor)
    r = capitulo.obtener(p, aid)
    assert r["usa_ia"] == 1 and r["aviso"] is None
    assert r["cifras_no_respaldadas"] == [{"seccion": "5.3", "cifra": "350"}]
    secciones = {s["numero"]: s["texto"] for s in r["capitulo"]["secciones"]}
    assert secciones["5.4"].startswith("Infiltración y lixiviación vertical")  # omitida por la IA → plantilla
    hechos = redactor.solicitudes[0]["messages"][0]["content"]
    assert "<hechos>" in hechos and '"envolvente":' not in hechos  # sin geometría

    doc = Document(io.BytesIO(capitulo.a_word(p, aid)))
    titulos = [x.text for x in doc.paragraphs if x.style.name.startswith("Heading")]
    assert titulos[0] == "5. ANÁLISIS DE DISPERSIÓN DE HIDROCARBUROS" and len(titulos) == 9
    assert len(doc.tables) == 3 and len(doc.inline_shapes) == 1
    marcadas = [run.text for x in doc.paragraphs for run in x.runs if run.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert marcadas == ["350"]


def test_capitulo_sin_ia_o_con_rechazo_usa_plantilla():
    p = crear()
    aid = capitulo.generar(p, usar_ia=True, cliente=RedactorFalso(stop="refusal"))
    r = capitulo.obtener(p, aid)
    assert r["usa_ia"] == 0 and "declinó" in r["aviso"]
    assert r["capitulo"]["secciones"] == r["capitulo"]["plantilla"]
    assert capitulo.obtener(p)["id"] == aid


def test_herramientas_del_asistente_y_tablero():
    from caepsa.asistente.herramientas import Contexto, ejecutar
    from caepsa.tablero import estado

    p = crear("Asistente dispersión", con_medio=False)
    ctx = Contexto(usuario="qa", proyecto=p.clave)
    avance, alertas = estado(p)
    assert avance[6] == 0 and any("Delimitación horizontal abierta" in a.mensaje for a in alertas)
    assert any(a.modulo == 6 and "sin valor o sin fuente" in a.mensaje for a in alertas)

    _, error = ejecutar(ctx, "actualizar_medio_fisico", {"datos": {"pendiente_pct": {"valor": 4, "fuente": ""}}})
    assert error  # la fuente es obligatoria
    salida, error = ejecutar(ctx, "actualizar_medio_fisico", {"datos": {
        "pendiente_pct": {"valor": 4, "fuente": "Levantamiento topográfico"},
        "direccion_pendiente": {"valor": "S", "fuente": "Levantamiento topográfico"}}})
    assert not error and json.loads(salida)["actualizados"] == ["pendiente_pct", "direccion_pendiente"]
    consulta = json.loads(ejecutar(ctx, "consultar_medio_fisico", {})[0])
    assert consulta["datos"]["pendiente_pct"]["valor"] == 4.0
    assert "Textura del suelo" in consulta["requeridos_sin_valor"]

    a = json.loads(ejecutar(ctx, "analizar_dispersion", {})[0])
    assert a["espacial"]["coincidencia_pendiente"]["direccion_pendiente"] == "S" and "envolvente" not in a["espacial"]

    salida, error = ejecutar(ctx, "generar_capitulo5", {"usar_ia": False})
    r = json.loads(salida)
    assert not error and r["redactado_con_ia"] is False and ctx.archivos_generados[-1].exists()
    avance, _ = estado(p)
    requeridos = sum(c.requerido for c in mf.CAMPOS)
    assert avance[6] == pytest.approx(0.6 * 2 / requeridos + 0.4)
